"""C3 tests on hand-built rate files whose problems are counted in advance."""
import csv
import gzip
import json

import pytest

from ptd_c3.cli import main
from ptd_c3.config import load_config
from ptd_c3.metrics import grade_scan
from ptd_c3.scan import npi_is_valid, scan_file
from ptd_c3.tags import load_tag_table

AS_OF = "2026-10-17"


def npi(base9: int) -> int:
    """Append the NPI check digit to a 9-digit base."""
    s = str(base9)
    total = 0
    for i, ch in enumerate(reversed("80840" + s)):
        d = int(ch)
        if i % 2 == 0:
            d *= 2
            d = d - 9 if d > 9 else d
        total += d
    return int(s + str((10 - total % 10) % 10))


def price(rate, ntype="negotiated", exp="9999-12-31", bclass="professional", **kw):
    p = {"negotiated_type": ntype, "negotiated_rate": rate, "expiration_date": exp, "billing_class": bclass}
    if bclass == "professional":
        p["service_code"] = ["11"]
    p.update(kw)
    return p


def item(code, prices, refs=(1,), code_type="CPT", arrangement="ffs"):
    return {"negotiation_arrangement": arrangement, "name": f"code {code}", "billing_code_type": code_type,
            "billing_code_type_version": "2026", "billing_code": code, "description": "x",
            "negotiated_rates": [{"provider_references": list(refs), "negotiated_prices": prices}]}


def refs(n=3, network="PPO - Large Group/Small Group"):
    return [{"provider_group_id": i + 1, "network_name": [network],
             "provider_groups": [{"npi": [npi(123456780 + i)], "tin": {"type": "ein", "value": str(950000000 + i)}}]}
            for i in range(n)]


def rate_file(items, references, refs_first=True, last_updated="2026-09-30"):
    head = {"reporting_entity_name": "Centene Management Company LLC", "reporting_entity_type": "TPA",
            "last_updated_on": last_updated, "version": "2.0.0"}
    body = [("provider_references", references), ("in_network", items)]
    return {**head, **dict(body if refs_first else body[::-1])}


@pytest.fixture(scope="module")
def cfg():
    return load_config(as_of=AS_OF)


@pytest.fixture(scope="module")
def files(tmp_path_factory):
    d = tmp_path_factory.mktemp("c3")
    clean_items = [item(str(99200 + i), [price(100.0 + i), price(150.0 + i, bclass="institutional")])
                   for i in range(10)]
    (d / "clean.json.gz").write_bytes(gzip.compress(json.dumps(rate_file(clean_items, refs())).encode()))

    # Dirty file: 20 price rows in total, problems counted by hand.
    dirty_refs = refs(2, network="") + [
        {"provider_group_id": 3, "network_name": [""], "provider_groups": [{"npi": [0], "tin": {"type": "ein", "value": "999999999"}}]},
        {"provider_group_id": 4, "network_name": [""], "provider_groups": [{"npi": [1234567890], "tin": {"type": "ein", "value": "950000009"}}]},
    ]
    outlier_prices = [price(v) for v in (100, 101, 99, 102, 98, 5000)]          # 6 rows, 1 outlier
    dirty_items = [
        item("99213", outlier_prices),
        item("99214", [price(0), price(0), price(80.0, exp="2025-01-01")], refs=(3,)),   # 2 zero (1 dup), 1 expired, TIN-only
        item("99215", [price(70.0), price(70.0)]),                                         # 1 exact duplicate
        item("99216", [price(40.0, ntype="percentage"), price(250.0, ntype="percentage")]),  # 1 pct > 100
        item("0450", [price(500.0, ntype="per diem", bclass="institutional")], code_type="RC"),
        item("CSTM-00", [price(12.0)], code_type="CSTM-ALL"),                              # catch-all
        item("9921", [price(60.0)]),                                                        # bad CPT format
        item("99217", [price(65.0)], refs=(4,)),                                            # invalid NPI only
        item("99218", [price(55.0)], refs=(77,)),                                           # dangling reference
        item("99219", [price(45.0, bclass="weird")]),                                       # bad enum
        item("99220", [{"negotiated_type": "negotiated", "negotiated_rate": 30.0, "billing_class": "institutional"}]),
    ]
    (d / "dirty.json").write_text(json.dumps(rate_file(dirty_items, dirty_refs, last_updated="2025-06-01")))

    # in_network before provider_references, refs 1-2 tagged PPO / HMO
    late_items = [item("99213", [price(100.0)], refs=(1,)), item("99214", [price(120.0)], refs=(2,)),
                  item("99215", [price(0)], refs=(2,))]
    (d / "late.json").write_text(json.dumps(rate_file(late_items, refs(2), refs_first=False)))

    # provider_references given by location
    (d / "loc_groups.json").write_text(json.dumps({"provider_groups": [
        {"npi": [npi(111111110)], "tin": {"type": "ein", "value": "951111111"}}]}))
    loc_refs = [{"provider_group_id": 1, "network_name": ["PPO"], "location": str(d / "loc_groups.json")}]
    (d / "loc.json").write_text(json.dumps(rate_file([item("99213", [price(90.0)])], loc_refs)))

    with open(d / "tags.csv", "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["npi", "tin", "tag"])
        w.writerow([npi(123456780), "", "PPO"])
        w.writerow(["", "950000001", "HMO"])
    with open(d / "expected.csv", "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["Specialty", "Billing Code"])
        for c, s in (("99213", "Family Medicine"), ("99214", "Family Medicine"), ("27447", "Orthopedics"),
                     ("99215", "Family Medicine")):
            w.writerow([s, c])
    with open(d / "c2_ppo_rate_files.csv", "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["rate_file_url", "ppo_variants", "network_label", "ppo_network_names", "needs_network_filter",
                    "ppo_plan_rows", "decided_by", "best_confidence", "confidence_counts"])
        w.writerow([str(d / "dirty.json"), "FULL", "UNCONFIRMED", "", "True", "3", "plan_name", "medium", "{}"])
        w.writerow([str(d / "clean.json.gz"), "FULL", "PPO_FULL", "PPO", "False", "5", "network_name", "high", "{}"])
    return d


def metrics_of(scan, cfg, expected=None, bucket="ALL"):
    _, ms = grade_scan(scan, cfg, expected)
    return {m["metric"]: m for m in ms if m["bucket"] == bucket}


# ------------------------------------------------------------------ unit

def test_npi_check_digit():
    assert npi_is_valid(1234567893)          # the CMS example NPI
    assert not npi_is_valid(1234567890)
    assert not npi_is_valid(123)
    assert npi_is_valid(npi(123456780))


def test_clean_file_scores_high(files, cfg):
    sc = scan_file(str(files / "clean.json.gz"), cfg)
    assert sc.status == "ok" and sc.acc["ALL"].rows == 20
    rows, _ = grade_scan(sc, cfg, None)
    assert rows[0]["score"] >= 95 and rows[0]["grade"] == "A"
    # no code list given, and no group has 5+ rates to test for outliers
    assert rows[0]["metrics_not_measured"] == "code_coverage,outliers"
    m = metrics_of(sc, cfg)
    assert m["provider_coverage"]["score"] == 1.0
    assert m["duplicates"]["score"] == 1.0


def test_dirty_file_counts(files, cfg):
    sc = scan_file(str(files / "dirty.json"), cfg)
    a = sc.acc["ALL"]
    assert a.rows == 20
    assert a.zero == 2
    assert a.expired == 1
    assert a.dups == 2                          # the second $0 row and the second 70.0 row
    assert a.catch_all == 1
    assert a.outliers == 2                      # 5000 among ~100s, and the 250% percentage
    assert a.neg_types["percentage"] == 2 and a.neg_types["per diem"] == 1
    assert a.groups_tin_only == 1 and a.groups_placeholder_tin == 1 and a.groups_invalid_npi == 1
    # rows tied to a group with a valid NPI: all but the 3 on ref 3, 1 on ref 4, 1 on dangling ref 77
    assert a.resolved == 15
    e = sc.schema_errors
    assert e["network_name_blank"] == 4
    assert e["billing_code_format"] == 1
    assert e["dangling_provider_reference"] == 1
    assert e["billing_class_invalid"] == 1
    assert e["price_field_missing"] == 1
    assert sc.examples["outlier"] and sc.examples["junk:tin_only_group"]


def test_dirty_file_scores_lower(files, cfg):
    clean = grade_scan(scan_file(str(files / "clean.json.gz"), cfg), cfg, None)[0][0]["score"]
    dirty_scan = scan_file(str(files / "dirty.json"), cfg)
    dirty = grade_scan(dirty_scan, cfg, None)[0][0]
    assert dirty["score"] < clean - 20
    m = metrics_of(dirty_scan, cfg)
    assert m["zero_rates"]["score"] == 0.0              # 10% zero share hits the tolerance
    assert 0 < m["staleness"]["score"] < 1               # one expired row, file 503 days old -> age score 0
    assert "2025-06-01" in m["staleness"]["justification"]


def test_expected_codes_coverage(files, cfg):
    from ptd_c3.codes import load_expected_codes

    exp = load_expected_codes(files / "expected.csv")
    assert exp.codes == {"99213", "99214", "99215", "27447"} and not exp.weights
    m = metrics_of(scan_file(str(files / "dirty.json"), cfg), cfg, exp)
    assert m["code_coverage"]["score"] == 0.75 and m["code_coverage"]["measured"]


def test_in_network_before_references_untagged(files, cfg):
    sc = scan_file(str(files / "late.json"), cfg)
    a = sc.acc["ALL"]
    assert sc.status == "ok" and a.rows == 3 and a.resolved == 3 and a.zero == 1


def test_provider_tags_split_buckets_even_when_references_come_last(files, cfg):
    tagger = load_tag_table(files / "tags.csv")
    sc = scan_file(str(files / "late.json"), cfg, tagger)
    assert sc.acc["ALL"].rows == 3
    assert sc.acc["PPO"].rows == 1 and sc.acc["PPO"].zero == 0
    assert sc.acc["HMO"].rows == 2 and sc.acc["HMO"].zero == 1
    assert sc.acc["ALL"].groups == 2                     # references not counted twice
    rows, _ = grade_scan(sc, cfg, None)
    assert [r["bucket"] for r in rows] == ["ALL", "HMO", "PPO"]


def test_location_references_are_fetched(files, cfg):
    sc = scan_file(str(files / "loc.json"), cfg)
    assert sc.location_fetched == 1 and sc.acc["ALL"].resolved == 1 and len(sc.acc["ALL"].npis) == 1


def test_byte_cap_marks_sampled(files, cfg):
    small = dict(cfg, stream=dict(cfg["stream"], max_bytes=600))
    sc = scan_file(str(files / "dirty.json"), small)
    assert sc.status == "sampled" and sc.bytes_read <= 600 + 65536


def test_missing_file_is_an_error_not_a_crash(files, cfg):
    sc = scan_file(str(files / "nope.json"), cfg)
    assert sc.status == "error"
    rows, _ = grade_scan(sc, cfg, None)
    assert rows[0]["status"] == "error"


# ------------------------------------------------------------------ CLI

def test_cli_reads_c2_list_and_writes_outputs(files, tmp_path):
    out = tmp_path / "c3"
    code = main(["run", "--c2", str(files / "c2_ppo_rate_files.csv"), "--as-of", AS_OF, "--out", str(out),
                 "--expected-codes", str(files / "expected.csv")])
    assert code == 0
    for f in ("c3_file_scores.csv", "c3_metrics.csv", "c3_schema_problems.csv", "c3_examples.csv",
              "c3_code_coverage.csv", "c3_results.xlsx", "c3_run_summary.json", "c3_run_log.txt"):
        assert (out / f).exists(), f
    with open(out / "c3_file_scores.csv") as fh:
        scores = {r["rate_file_url"].rsplit("/", 1)[-1]: r for r in csv.DictReader(fh)}
    assert float(scores["clean.json.gz"]["score"]) > float(scores["dirty.json"]["score"])
    assert scores["dirty.json"]["needs_network_filter"] == "True"     # carried over from C2
    with open(out / "c3_metrics.csv") as fh:
        ms = list(csv.DictReader(fh))
    assert len(ms) == 2 * 9 and all(m["justification"] for m in ms)
    log = (out / "c3_run_log.txt").read_text()
    assert "graded as a whole" in log


def test_cli_files_and_tags(files, tmp_path):
    out = tmp_path / "c3t"
    assert main(["run", "--files", str(files / "late.json"), "--provider-tags", str(files / "tags.csv"),
                 "--as-of", AS_OF, "--out", str(out)]) == 0
    with open(out / "c3_file_scores.csv") as fh:
        assert {r["bucket"] for r in csv.DictReader(fh)} == {"ALL", "PPO", "HMO"}


def test_cli_missing_c2_list(tmp_path):
    assert main(["run", "--c2", str(tmp_path / "none.csv"), "--out", str(tmp_path)]) == 1


def test_hilabs_weights_workbook_weighted_coverage(files, cfg, tmp_path):
    """Same sheet layout as the HiLabs 'Rate Benchmarking Individual Weights' workbook, saved without an extension."""
    import openpyxl

    from ptd_c3.codes import load_expected_codes

    wb = openpyxl.Workbook()
    wb.active.title = "Taxonomy-Spec Xwalk"
    wb.active.append(["SPEC_TYPE", "SPEC_CD", "SPEC_DESCR", "Taxonomy Code"])
    wb.active.append(["Individual", "PCP", "Primary Care", "207Q00000X"])
    sw = wb.create_sheet("Specialty Weights")
    sw.append(["Spec Code", "Specialty", "Specialty Weight"])
    sw.append(["PCP", "Primary Care", 300])
    sw.append(["025", "Orthopedic Surgery", 100])
    sp = wb.create_sheet("Spec-Proc Weights")
    sp.append(["SPEC_CD", "SPEC_DESCRIPTION", "PROC_CD", "SPEC_PROC_RANK", "SPEC_PROC_WEIGHT"])
    sp.append(["PCP", "Primary Care", "99213", 2, 10])
    sp.append(["PCP", "Primary Care", "27447", 1, 20])
    sp.append(["025", "Orthopedic Surgery", "99214", 1, 7])
    sp.append(["080", "Urgent Care", "99215", 1, 5])         # facility: no specialty weight
    path = tmp_path / "weights_no_extension"
    wb.save(path)

    exp = load_expected_codes(path)
    assert exp.codes == {"99213", "27447", "99214", "99215"}
    assert exp.weights == {"99213": 100.0, "27447": 200.0, "99214": 100.0}
    assert exp.unweighted_specialties == ["080 Urgent Care"]
    m = metrics_of(scan_file(str(files / "dirty.json"), cfg), cfg, exp)["code_coverage"]
    assert m["score"] == 0.5                                  # 99213 + 99214 = 200 of 400, while 3 of 4 codes are present
    assert "75.00%" in m["justification"] and "50.00%" in m["justification"]

    out = tmp_path / "c3w"
    assert main(["run", "--files", str(files / "dirty.json"), "--expected-codes", str(path), "--as-of", AS_OF,
                 "--out", str(out)]) == 0
    with open(out / "c3_code_coverage.csv") as fh:
        rows = {r["specialty"]: r for r in csv.DictReader(fh)}
    assert rows["PCP Primary Care"]["coverage"] == "0.5" and rows["PCP Primary Care"]["weighted_coverage"] == "0.3333"
    assert "Urgent Care" in (out / "c3_run_log.txt").read_text()
