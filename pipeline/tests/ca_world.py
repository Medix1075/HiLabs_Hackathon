"""A miniature California world for C4-C7, built so that every right answer is known in advance.

One Health Net-style rate file mixes PPO and HMO provider references with a blank network_name
(as the real CA file does). Around it: an NPPES extract (main file + practice locations), a CMS
Doctors & Clinicians file, a payer directory and a HiLabs-layout workbook.

Ground truth kept in `truth`:
  ref_network   provider_group_id -> PPO / HMO
  tin_site      TIN -> its real practice site
  tier          provider_group_id -> price factor (0.8 cheap ... 1.4 expensive)
"""
from __future__ import annotations

import csv
import gzip
import json
import random
from pathlib import Path

PPO_CODES = {"99213": 100.0, "99214": 140.0, "99203": 120.0, "93000": 25.0, "93306": 220.0}
SITES = [  # (line1, line2, city, zip)
    ("1200 West Third Street", "Suite 100", "Los Angeles", "90017"),
    ("455 N Brand Blvd", "", "Glendale", "91203"),
    ("2100 Webster St", "Ste 200", "San Francisco", "94115"),
    ("3900 Fifth Avenue", "", "San Diego", "92103"),
    ("1 Medical Plaza Dr", "", "Irvine", "92618"),
    ("700 E Shaw Ave", "", "Fresno", "93710"),
]
TAX_PCP, TAX_CARD = "207R00000X", "207RC0000X"


def npi(base9: int) -> int:
    s = str(base9)
    total = 0
    for i, ch in enumerate(reversed("80840" + s)):
        d = int(ch)
        if i % 2 == 0:
            d *= 2
            d = d - 9 if d > 9 else d
        total += d
    return int(s + str((10 - total % 10) % 10))


def price(rate, ntype="negotiated", exp="9999-12-31", bclass="professional", pos=("11",), mods=None):
    p = {"negotiated_type": ntype, "negotiated_rate": round(rate, 2), "expiration_date": exp, "billing_class": bclass}
    if bclass == "professional":
        p["service_code"] = list(pos)
    if mods:
        p["billing_code_modifier"] = list(mods)
    return p


def build(root: Path) -> dict:
    rnd = random.Random(11)
    root.mkdir(parents=True, exist_ok=True)
    refs, truth = [], {"ref_network": {}, "tin_site": {}, "tier": {}, "tin_npis": {}, "npi_tax": {}}
    nxt = [100000000]

    def new_npi():
        nxt[0] += 1
        return npi(nxt[0])

    # 40 PPO references and 30 HMO references, one TIN each, 3 individual NPIs + 1 organisation NPI
    def make_ref(gid, network, tin, tier):
        members = [new_npi() for _ in range(3)]
        org = new_npi()
        truth["ref_network"][gid], truth["tier"][gid] = network, tier
        truth["tin_npis"].setdefault(tin, set()).update(members)
        truth.setdefault("org_npi", {})[tin] = org
        truth["tin_site"].setdefault(tin, SITES[int(tin) % len(SITES)])
        for i, m in enumerate(members):
            truth["npi_tax"][m] = TAX_CARD if (int(tin) % 4 == 0) else TAX_PCP
        refs.append({"provider_group_id": gid, "network_name": [""],
                     "provider_groups": [{"npi": members + [org], "tin": {"type": "ein", "value": tin}}]})
        return members

    tiers = [0.8, 1.0, 1.2, 1.4]
    for g in range(1, 41):
        make_ref(g, "PPO", str(950000000 + g), tiers[(g - 1) // 10])
    for g in range(41, 71):
        make_ref(g, "HMO", str(950000000 + g), 0.75)
    # 20 dual-contract TINs: the TIN of PPO refs 1-20 also has an HMO reference (71-90) with the same NPIs
    for k, g in enumerate(range(71, 91)):
        ppo_ref = refs[k]
        tin = ppo_ref["provider_groups"][0]["tin"]["value"]
        truth["ref_network"][g], truth["tier"][g] = "HMO", 0.75
        refs.append({"provider_group_id": g, "network_name": [""],
                     "provider_groups": [{"npi": list(ppo_ref["provider_groups"][0]["npi"]),
                                          "tin": {"type": "ein", "value": tin}}]})
    # One TIN-only group (no NPI): must be counted, not extracted
    refs.append({"provider_group_id": 99, "network_name": [""],
                 "provider_groups": [{"npi": [], "tin": {"type": "ein", "value": "950000099"}}]})
    truth["ref_network"][99] = "PPO"

    ppo_ids = list(range(1, 41))
    hmo_ids = list(range(41, 91))
    items = []
    for code, base in PPO_CODES.items():
        nrs = []
        # PPO fee schedules: refs co-listed in blocks of 5 that share one tier
        for start in range(0, 40, 5):
            block = ppo_ids[start:start + 5]
            factor = truth["tier"][block[0]]
            nrs.append({"provider_references": block, "negotiated_prices": [price(base * factor)]})
        # HMO schedules: blocks of 6 (refs 41-90), cheaper
        for start in range(0, len(hmo_ids), 6):
            block = hmo_ids[start:start + 6]
            nrs.append({"provider_references": block, "negotiated_prices": [price(base * 0.75)]})
        items.append({"negotiation_arrangement": "ffs", "name": code, "billing_code_type": "CPT",
                      "billing_code_type_version": "2026", "billing_code": code, "description": f"desc {code}",
                      "negotiated_rates": nrs})
    # Variants and junk on 99213 for ref 1's block: institutional, modifier, zero, expired, duplicate, percentage
    extra = [{"provider_references": [1, 2, 3, 4, 5], "negotiated_prices": [
        price(300.0, bclass="institutional"), price(90.0, mods=["25"]), price(0.0), price(55.0, exp="2024-12-31"),
        price(80.0), price(150.0, ntype="percentage")]},
        # ref 6 gets a second, competing rate for the same variant: tie-break by support (80 x2) vs 85
        {"provider_references": [6], "negotiated_prices": [price(85.0)]}]
    items[0]["negotiated_rates"].extend(extra)
    # Capitation for half of the HMO references (an HMO arrangement): network evidence
    cap_refs = [r for r in hmo_ids if r % 2 == 1]
    items.append({"negotiation_arrangement": "capitation", "name": "cap", "billing_code_type": "CPT",
                  "billing_code_type_version": "2026", "billing_code": "99499", "description": "capitation",
                  "negotiated_rates": [{"provider_references": cap_refs,
                                        "negotiated_prices": [price(30.0, ntype="fee schedule")]}]})
    # An out-of-scope code (not in the workbook): read for evidence, not extracted
    items.append({"negotiation_arrangement": "ffs", "name": "x", "billing_code_type": "CPT",
                  "billing_code_type_version": "2026", "billing_code": "11111", "description": "out of scope",
                  "negotiated_rates": [{"provider_references": [1], "negotiated_prices": [price(10.0)]}]})
    rate_file = {"reporting_entity_name": "Centene Management Company LLC", "reporting_entity_type": "TPA",
                 "last_updated_on": "2026-09-30", "version": "2.0.0", "provider_references": refs, "in_network": items}
    rf = root / "healthnet-ca_in-network.json.gz"
    rf.write_bytes(gzip.compress(json.dumps(rate_file).encode()))
    # The same file with in_network first, to prove order does not matter
    late = {k: rate_file[k] for k in ("reporting_entity_name", "last_updated_on", "version", "in_network",
                                      "provider_references")}
    (root / "late_in-network.json").write_text(json.dumps(late))

    # ------------------------------------------------------------ directory: some PPO and HMO NPIs
    with open(root / "directory.csv", "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["npi", "tin", "network_name", "address_line1", "address_line2", "city", "state", "zip"])
        for g in range(1, 36):  # the directory covers most, not all, of the PPO network
            grp = refs[g - 1]["provider_groups"][0]
            site = truth["tin_site"][grp["tin"]["value"]]
            for m in grp["npi"][:3]:
                w.writerow([m, grp["tin"]["value"], "PPO - Large Group/Small Group", site[0], site[1], site[2], "CA",
                            site[3]])
        # a listing under the payer's name only (no network named): a candidate, never a confirmation
        w.writerow([refs[35]["provider_groups"][0]["npi"][0], "", "Health Net", "", "", "", "", ""])
        for g in range(41, 46):
            grp = refs[g - 1]["provider_groups"][0]
            for m in grp["npi"][:3]:
                w.writerow([m, "", "HMO - Full Network", "", "", "", "", ""])

    # ------------------------------------------------------------ NPPES main + practice locations
    special = {}
    first_ppo = refs[0]["provider_groups"][0]
    second_ppo = refs[19]["provider_groups"][0]      # tier 1.0 group
    moved = first_ppo["npi"][0]                      # also practises under TIN of ref 20 (its own site)
    special["two_tin_npi"] = moved
    refs[19]["provider_groups"][0]["npi"].append(moved)
    rf.write_bytes(gzip.compress(json.dumps(rate_file).encode()))
    (root / "late_in-network.json").write_text(json.dumps({**late, "provider_references": refs}))
    special["two_tin_tins"] = (first_ppo["tin"]["value"], second_ppo["tin"]["value"])
    stale_npi = refs[24]["provider_groups"][0]["npi"][1]
    special["stale_npi"], special["stale_tin"] = stale_npi, refs[24]["provider_groups"][0]["tin"]["value"]
    deact_npi = refs[30]["provider_groups"][0]["npi"][2]
    special["deactivated_npi"] = deact_npi
    no_nppes = refs[33]["provider_groups"][0]["npi"][0]
    special["not_in_nppes"] = no_nppes

    header = ["NPI", "Entity Type Code", "Provider Organization Name (Legal Business Name)",
              "Provider Last Name (Legal Name)", "Provider First Name",
              "Provider First Line Business Mailing Address", "Provider Second Line Business Mailing Address",
              "Provider Business Mailing Address City Name", "Provider Business Mailing Address State Name",
              "Provider Business Mailing Address Postal Code",
              "Provider First Line Business Practice Location Address",
              "Provider Second Line Business Practice Location Address",
              "Provider Business Practice Location Address City Name",
              "Provider Business Practice Location Address State Name",
              "Provider Business Practice Location Address Postal Code", "Last Update Date",
              "NPI Deactivation Date", "NPI Reactivation Date",
              "Healthcare Provider Taxonomy Code_1", "Healthcare Provider Primary Taxonomy Switch_1",
              "Healthcare Provider Taxonomy Code_2", "Healthcare Provider Primary Taxonomy Switch_2"]
    pl_rows = []
    with open(root / "npidata_pfile.csv", "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(header)
        w.writerow(["1999999999", "1", "", "Nobody", "Not", "", "", "", "", "", "1 A St", "", "X", "CA", "90001",
                    "01/01/2020", "", "", TAX_PCP, "Y", "", ""])  # an NPI not in the rate file
        done = set()
        for ref in refs:
            grp = ref["provider_groups"][0]
            tin = grp["tin"]["value"]
            if tin not in truth["tin_site"]:
                continue
            site = truth["tin_site"][tin]
            for m in grp["npi"]:
                if m in done or m == no_nppes:
                    continue
                done.add(m)
                is_org = m == truth["org_npi"].get(tin)
                if m == stale_npi:
                    addr, upd = (f"{m % 997} Old Mill Road", "", "Bakersfield", "93301"), "03/15/2012"
                elif is_org or rnd.random() < 0.6:
                    # vary the spelling so standardisation has work to do
                    addr = (site[0].upper().replace("STREET", "ST").replace("AVENUE", "AVE"), site[1].replace("Suite", "STE"),
                            site[2], site[3] + "1234")
                    upd = "06/01/2025"
                else:  # a stale address of its own (where it worked years ago)
                    addr, upd = (f"{m % 997} Old Mill Road", "", "Bakersfield", "93301"), "01/10/2016"
                deact = "05/05/2024" if m == deact_npi else ""
                tax = truth["npi_tax"].get(m, "193200000X")
                w.writerow([m, "2" if is_org else "1", f"Group {tin}" if is_org else "", "" if is_org else f"Doe{m % 1000}",
                            "" if is_org else "Pat", "PO Box 100", "", "Sacramento", "CA", "95814",
                            addr[0], addr[1], addr[2], "CA", addr[3], upd, deact, "", tax, "Y", "", ""])
                if m == stale_npi:
                    pl_rows.append([m, site[0], site[1], site[2], "CA", site[3]])
    with open(root / "pl_pfile.csv", "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["NPI", "Provider Secondary Practice Location Address- Address Line 1",
                    "Provider Secondary Practice Location Address-  Address Line 2",
                    "Provider Secondary Practice Location Address - City Name",
                    "Provider Secondary Practice Location Address - State Name",
                    "Provider Secondary Practice Location Address - Postal Code"])
        w.writerows(pl_rows)

    # ------------------------------------------------------------ CMS Doctors and Clinicians
    with open(root / "DAC_NationalDownloadableFile.csv", "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["NPI", "Ind_PAC_ID", "Provider Last Name", "Provider First Name", "pri_spec", "Facility Name",
                    "org_pac_id", "num_org_mem", "adr_ln_1", "adr_ln_2", "ln_2_sprs", "City/Town", "State",
                    "ZIP Code", "adrs_id"])
        for tin, members in truth["tin_npis"].items():
            site = truth["tin_site"][tin]
            for m in sorted(members):
                if rnd.random() < 0.5 or m == moved:
                    w.writerow([m, "1", "Doe", "Pat", "INTERNAL MEDICINE", f"Group {tin}", f"P{tin}", "4",
                                site[0], site[1], "", site[2], "CA", site[3], "x"])
        # the two-TIN NPI also works for the second TIN's group, at that group's site
        t2 = special["two_tin_tins"][1]
        s2 = truth["tin_site"][t2]
        w.writerow([moved, "1", "Doe", "Pat", "INTERNAL MEDICINE", f"Group {t2}", f"P{t2}", "4",
                    s2[0], s2[1], "", s2[2], "CA", s2[3], "y"])
        for m in sorted(truth["tin_npis"][t2]):
            w.writerow([m, "1", "Doe", "Pat", "INTERNAL MEDICINE", f"Group {t2}", f"P{t2}", "4",
                        s2[0], s2[1], "", s2[2], "CA", s2[3], "y"])

    # ------------------------------------------------------------ HiLabs-layout workbook
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Taxonomy-Spec Xwalk"
    ws.append(["SPEC_TYPE", "SPEC_CD", "SPEC_DESCR", "Taxonomy Code"])
    ws.append(["Individual", "PCP", "Primary Care", TAX_PCP])
    ws.append(["Individual", "010", "Cardiology", TAX_CARD])
    ws = wb.create_sheet("Specialty Weights")
    ws.append(["Spec Code", "Specialty", "Specialty Weight"])
    ws.append(["PCP", "Primary Care", 1000])
    ws.append(["010", "Cardiology", 300])
    ws = wb.create_sheet("Spec-Proc Weights")
    ws.append(["SPEC_CD", "SPEC_DESCRIPTION", "PROC_CD", "SPEC_PROC_RANK", "SPEC_PROC_WEIGHT"])
    for rank, (code, wt) in enumerate([("99213", 100), ("99214", 80), ("99203", 30)], 1):
        ws.append(["PCP", "Primary Care", code, rank, wt])
    for rank, (code, wt) in enumerate([("93000", 50), ("93306", 40), ("99214", 20)], 1):
        ws.append(["010", "Cardiology", code, rank, wt])
    wb.save(root / "weights.xlsx")

    with open(root / "zip_county.csv", "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["zip", "county"])
        for s, c in zip(SITES, ["Los Angeles", "Los Angeles", "San Francisco", "San Diego", "Orange", "Fresno"]):
            w.writerow([s[3], c + " County"])

    return {"root": root, "rate_file": rf, "late_file": root / "late_in-network.json",
            "directory": root / "directory.csv", "nppes": root / "npidata_pfile.csv", "pl": root / "pl_pfile.csv",
            "dac": root / "DAC_NationalDownloadableFile.csv", "workbook": root / "weights.xlsx",
            "zip_county": root / "zip_county.csv", "truth": truth, "special": special}
