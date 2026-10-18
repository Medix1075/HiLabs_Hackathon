"""C4-C7 tests on the miniature California world in ca_world.py, where every right answer is known."""
import csv
import gzip
import http.server
import json
import sqlite3
import threading
import zipfile
from collections import defaultdict
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import numpy as np
import openpyxl
import pytest

import ca_world
from ptd_c3.tags import load_tag_table
from ptd_c4.cli import main as c4
from ptd_c4.directory import crawl_plannet
from ptd_c5.address import standardize
from ptd_c5.cli import main as c5
from ptd_c5.fetch import extract_nppes
from ptd_c6.cli import main as c6
from ptd_c7.bench import group_stats
from ptd_c7.cli import main as c7
from ptd_run.cli import main as run_all

AS_OF = "2026-10-17"


@pytest.fixture(scope="module")
def world(tmp_path_factory):
    root = tmp_path_factory.mktemp("ca")
    w = ca_world.build(root / "inputs")
    out = root / "out"
    db = str(out / "rates.sqlite")
    assert c4(["run", "--files", str(w["rate_file"]), "--workbook", str(w["workbook"]), "--directory",
               str(w["directory"]), "--db", db, "--out", str(out / "c4"), "--workers", "2", "--as-of", AS_OF]) == 0
    assert c5(["run", "--db", db, "--nppes", str(w["nppes"]), "--nppes-pl", str(w["pl"]), "--dac", str(w["dac"]),
               "--directory", str(w["directory"]), "--zip-county", str(w["zip_county"]), "--out", str(out / "c5"),
               "--as-of", AS_OF]) == 0
    assert c6(["run", "--db", db, "--workbook", str(w["workbook"]), "--out", str(out / "c6")]) == 0
    assert c7(["run", "--db", db, "--workbook", str(w["workbook"]), "--c6-out", str(out / "c6"),
               "--out", str(out / "c7")]) == 0
    w.update(out=out, db=db, con=sqlite3.connect(db))
    yield w
    w["con"].close()


def rows(con, q, *a):
    return con.execute(q, a).fetchall()


# ------------------------------------------------------------------ C4: network tags
def test_no_reference_is_tagged_with_the_wrong_network(world):
    truth = world["truth"]["ref_network"]
    wrong = [(ref, tag) for ref, tag in rows(world["con"], "SELECT ref_id, tag FROM c4_refs")
             if (truth[int(ref)] == "PPO" and tag == "NON_PPO") or (truth[int(ref)] == "HMO" and tag == "PPO")]
    assert wrong == []
    tags = dict(rows(world["con"], "SELECT tag, COUNT(*) FROM c4_refs GROUP BY tag"))
    assert tags["NON_PPO"] == 50          # every HMO reference found, from capitation, exclusion and spreading
    assert tags["PPO"] >= 20


def test_hypotheses_are_measured_before_use(world):
    s = json.loads((world["out"] / "c4/c4_run_summary.json").read_text())
    rep = next(iter(s["network"].values()))
    assert rep["propagation"]["used"] and rep["propagation"]["holdout_accuracy"] >= 0.9
    assert rep["dual_contract"]["precision"] == 1.0 and rep["dual_contract"]["tested"] >= 10
    assert rep["dual_contract"]["used"] is True


def test_shared_npi_directory_listing_goes_to_the_reference_not_excluded(world):
    ev = dict(rows(world["con"], "SELECT ref_id, evidence FROM c4_refs WHERE ref_id IN ('1', '71')"))
    assert "their other reference(s) 71 are NON_PPO" in ev["1"]
    assert "capitation" in ev["71"]


def test_network_status_is_never_wrong_when_confirmed(world):
    truth = world["truth"]["ref_network"]
    got = dict(rows(world["con"], "SELECT ref_id, status FROM c4_refs"))
    wrong = [(r, st) for r, st in got.items() if (st == "CONFIRMED_PPO" and truth[int(r)] != "PPO")
             or (st == "CONFIRMED_NON_PPO" and truth[int(r)] != "HMO")]
    assert wrong == []
    n = {st: sum(1 for v in got.values() if v == st) for st in set(got.values())}
    assert n["CONFIRMED_PPO"] >= 9 and n["CONFIRMED_NON_PPO"] >= 25
    # spreading and shared listings never confirm: they leave candidates / unresolved
    basis = dict(rows(world["con"], "SELECT ref_id, status_basis FROM c4_refs WHERE status = 'PPO_CANDIDATE'"))
    assert basis and all(b.startswith(("inferred", "listed")) for b in basis.values())


def test_hmo_references_are_not_extracted(world):
    hmo_tins = {str(950000000 + g) for g in range(41, 71)}
    confirmed_tins = {t for (t,) in rows(world["con"], "SELECT DISTINCT tin FROM c4_rates WHERE network_status = 'CONFIRMED_PPO'")}
    assert confirmed_tins and not (confirmed_tins & hmo_tins)
    non = {r for (r,) in rows(world["con"], "SELECT ref_id FROM c4_refs WHERE status = 'CONFIRMED_NON_PPO'")}
    extracted = {r for (r,) in rows(world["con"], "SELECT DISTINCT ref_id FROM c4_sig_members")}
    assert non and not (non & extracted)
    tins = {t for (t,) in rows(world["con"], "SELECT DISTINCT tin FROM c4_pairs")}
    assert "950000099" not in tins                            # TIN-only group: no NPI, no rows


def test_variants_are_split_and_junk_dropped(world):
    con, npi = world["con"], world["truth"]["tin_npis"]["950000002"]
    n = sorted(npi)[0]
    got = {(r[0], r[1], r[2]): r[3] for r in rows(
        con, "SELECT billing_class, negotiated_type, modifier, rate FROM c4_rates WHERE npi = ? AND code = '99213'", n)}
    assert got[("professional", "negotiated", "")] == 80.0    # zero and expired rows dropped, repeat merged
    assert got[("professional", "negotiated", "25")] == 90.0
    assert got[("institutional", "negotiated", "")] == 300.0
    assert got[("professional", "percentage", "")] == 150.0


def test_competing_rates_use_the_stated_tie_break(world):
    n = sorted(world["truth"]["tin_npis"]["950000006"])[0]
    r = rows(world["con"], "SELECT rate, distinct_rates, pick_rule FROM c4_rates WHERE npi = ? AND code = '99213' "
             "AND billing_class = 'professional' AND modifier = '' AND negotiated_type = 'negotiated'", n)
    assert r == [(80.0, 2, "lowest")]


def test_code_scope_and_capitation(world):
    codes = {c for (c,) in rows(world["con"], "SELECT DISTINCT code FROM c4_sig_rates")}
    assert codes == set(ca_world.PPO_CODES)                    # 11111 and 99499 are outside the workbook


def test_reference_order_and_worker_count_do_not_change_rates(world, tmp_path):
    def rates(path, workers, codes="all"):
        db = tmp_path / f"r{workers}.sqlite"
        assert c4(["run", "--files", str(path), "--codes", codes, "--directory", str(world["directory"]), "--db", str(db),
                   "--out", str(tmp_path / f"o{workers}"), "--workers", str(workers), "--as-of", AS_OF]) == 0
        con = sqlite3.connect(db)
        out = sorted(con.execute("SELECT npi, tin, code, negotiated_type, billing_class, modifier, rate FROM c4_rates"))
        con.close()
        return out
    a = rates(world["rate_file"], 1)
    b = rates(world["late_file"], 3)
    assert a == b and len(a) > 0
    assert not [r for r in a if r[2] == "99499"]               # capitation rows are never extracted


def test_provider_tags_feed_c3(world):
    t = load_tag_table(world["out"] / "c4/c4_provider_tags.csv")
    assert t.by_tin["950000041"] == {"CONFIRMED_NON_PPO"}
    assert "CONFIRMED_PPO" in t.by_tin["950000001"]


# ------------------------------------------------------------------ C5: addresses
def test_every_resolved_pair_gets_its_tins_real_site(world):
    truth = world["truth"]["tin_site"]
    wrong = []
    for npi, tin, site in rows(world["con"], "SELECT npi, tin, site_key FROM c5_addresses WHERE rank = 1 AND address != ''"):
        t = truth[tin]
        if standardize(t[0], t[1], t[2], "CA", t[3]).site_key != site:
            wrong.append((npi, tin))
    assert wrong == []


def test_one_npi_two_tins_two_addresses(world):
    sp = world["special"]
    got = dict(rows(world["con"], "SELECT tin, address FROM c5_addresses WHERE npi = ? AND rank = 1", sp["two_tin_npi"]))
    a, b = sp["two_tin_tins"]
    assert got[a] != got[b]
    assert "IRVINE" in got[b] and "SAN DIEGO" in got[a]


def test_stale_deactivated_and_missing_npis(world):
    sp, con = world["special"], world["con"]
    stale = rows(con, "SELECT address, verified FROM c5_addresses WHERE npi = ? AND rank = 1", sp["stale_npi"])[0]
    assert "OLD MILL" not in stale[0] and stale[1]
    assert rows(con, "SELECT address, evidence FROM c5_addresses WHERE npi = ?", sp["deactivated_npi"]) == \
        [("", "NPI deactivated in NPPES")]
    missing = rows(con, "SELECT address, families FROM c5_addresses WHERE npi = ? AND rank = 1", sp["not_in_nppes"])[0]
    assert missing[0] and "tin_consensus" in missing[1]


def test_logic_file_and_stats(world):
    text = (world["out"] / "c5/c5_logic.md").read_text()
    for part in ("TIN consensus", "Group alignment", "Conflict rules", "Verified address", "This run"):
        assert part in text
    s = json.loads((world["out"] / "c5/c5_run_summary.json").read_text())
    assert s["pct_npis_with_address"] > 99 and s["pct_npis_with_verified_address"] > 90
    assert s["best_address_vs_nppes_practice"]["differs from NPPES"] > 0


def test_standardiser():
    a = standardize("1200 West Third Street, Suite 100", "", "Los Angeles", "ca", "900171234")
    b = standardize("Medical Plaza", "1200 W. 3rd St. #100", "LOS ANGELES", "CA", "90017")
    assert a.site_key == b.site_key == "1200|W 3RD ST|90017"
    assert a.unit == b.unit == "STE 100"
    assert standardize("P.O. Box 55", "", "Fresno", "CA", "93701").po_box


def test_c5_without_sources_leaves_pairs_unresolved(world, tmp_path):
    import shutil
    db = tmp_path / "r.sqlite"
    shutil.copy(world["db"], db)
    assert c5(["run", "--db", str(db), "--out", str(tmp_path / "c5"), "--as-of", AS_OF]) == 0
    s = json.loads((tmp_path / "c5/c5_run_summary.json").read_text())
    assert s["pct_npis_with_address"] == 0


# ------------------------------------------------------------------ C6: the rate file
def test_rate_file_columns_rows_and_headline(world):
    wb = openpyxl.load_workbook(world["out"] / "c6/c6_provider_rate_file.xlsx", read_only=True)
    ws = wb["Rates"]
    data = list(ws.iter_rows(values_only=True))
    assert list(data[0][:6]) == ["NPI", "TIN", "Address", "Billing Code", "Specialty", "Rate"]
    s = json.loads((world["out"] / "c6/c6_run_summary.json").read_text())
    assert len(data) - 1 == s["files"]["c6_provider_rate_file"]["rows"] > 0
    h = data[0]
    assert {r[h.index("Network Status")] for r in data[1:]} == {"CONFIRMED_PPO"}
    by = defaultdict(list)
    for r in data[1:]:
        by[(r[h.index("NPI")], r[h.index("Billing Code")])].append(r)
    org = world["truth"]["org_npi"]["950000002"]
    person = sorted(world["truth"]["tin_npis"]["950000002"])[0]
    assert by[(person, "99213")][0][h.index("Rate")] == 80.0     # Type 1: professional, no modifier, office
    assert by[(person, "99213")][0][h.index("Price Variants for Code")] == 4
    assert by[(person, "99213")][0][h.index("Network Basis")].startswith("direct: own directory listing")
    # The group's organisation NPI is not listed itself: its reference's listed doctors don't confirm it,
    # so it is a candidate (with the same contract's rates).
    assert (org, "99213") not in by
    cw = openpyxl.load_workbook(world["out"] / "c6/c6_provider_rate_file_candidates.xlsx", read_only=True)["Rates"]
    cd = list(cw.iter_rows(values_only=True))
    ch = cd[0]
    org_rows = [r for r in cd[1:] if r[ch.index("NPI")] == org and r[ch.index("Billing Code")] == "99213"]
    assert org_rows[0][ch.index("Rate")] == 300.0                 # Type 2: institutional when it exists
    assert org_rows[0][ch.index("Network Status")] == "PPO_CANDIDATE"
    assert org_rows[0][ch.index("Network Basis")].startswith("inferred: other members of reference")


def test_specialty_sources_and_plausibility(world):
    prov = {int(r["npi"]): r for r in csv.DictReader(open(world["out"] / "c6/c6_providers.csv"))}
    pcp = next(n for n, t in world["truth"]["npi_tax"].items() if t == ca_world.TAX_PCP and n in prov)
    assert prov[pcp]["specialty"] == "Primary Care" and prov[pcp]["specialty_source"] == "taxonomy_primary"
    assert prov[world["special"]["not_in_nppes"]]["specialty_source"] == "inferred_from_codes"
    exp = dict(rows(world["con"], "SELECT code, code_expected FROM c6_rates WHERE npi = ?", pcp))
    assert exp["99213"] == 1 and exp["93000"] == 0


# ------------------------------------------------------------------ C7: cost effectiveness
def test_cost_index_follows_the_true_price_tiers(world):
    idx = defaultdict(list)
    for r in csv.DictReader(open(world["out"] / "c7/c7_provider_scores.csv")):
        g = int(r["tin"]) - 950000000
        if g <= 40:
            idx[world["truth"]["tier"][g]].append(float(r["cost_index"]))
    means = [np.mean(idx[t]) for t in (0.8, 1.0, 1.2, 1.4)]
    assert means == sorted(means) and means[0] < 0.9 < 1.1 < means[3]


def test_group_stats_quantiles_and_tied_ranks():
    gid = np.array([0, 0, 0, 0, 1, 1])
    vals = np.array([10.0, 20.0, 20.0, 40.0, 5.0, 7.0])
    size, q1, med, q3, pct = group_stats(gid, vals, 2)
    assert list(size) == [4, 4, 4, 4, 2, 2]
    assert med[0] == 20.0 and q1[0] == 17.5 and q3[0] == 25.0 and med[4] == 6.0
    assert list(pct[:4]) == [0.0, 0.5, 0.5, 1.0]


def test_metrics_file_and_scored_rate_file(world):
    text = (world["out"] / "c7/c7_metrics.md").read_text()
    for part in ("similar providers", "Rate index", "shrinkage", "This run"):
        assert part in text
    s7 = json.loads((world["out"] / "c7/c7_run_summary.json").read_text())
    s6 = json.loads((world["out"] / "c6/c6_run_summary.json").read_text())
    assert s7["scored_rate_file_rows"] == s6["files"]["c6_provider_rate_file"]["rows"]
    assert s7["set"] == "confirmed" and s7["network_statuses"] == ["CONFIRMED_PPO"]
    with gzip.open(world["out"] / "c7/c7_provider_rate_file_scored.csv.gz", "rt") as fh:
        head = next(csv.reader(fh))
    assert head[-4:] == ["Provider Cost Index", "Provider Cost Score", "Provider Tier", "Score Confidence"]


# ------------------------------------------------------------------ the runner and the summary
def test_runner_writes_summary_and_deliverables(world, tmp_path):
    w = world
    code = run_all(["all", "--from", "c3", "--files", str(w["rate_file"]), "--workbook", str(w["workbook"]),
                    "--directory", str(w["directory"]), "--nppes", str(w["nppes"]), "--nppes-pl", str(w["pl"]),
                    "--dac", str(w["dac"]), "--workers", "2", "--as-of", AS_OF, "--grade-by-network",
                    "--out", str(tmp_path / "out"), "--data", str(tmp_path / "data")])
    assert code == 0
    text = (tmp_path / "out/summary.md").read_text()
    assert "% NPIs linked to a verified address (C5)" in text and "| Unique NPIs (C4) | 212 |" in text
    assert "| Network status of provider references (C4) | CONFIRMED_PPO: 35;" in text
    # 35 listed PPO references, 3 of 4 members listed: the unlisted organisation NPIs are candidates
    assert "| Network status of NPI-TIN pairs, decided per pair (C4) | CONFIRMED_PPO: 105; PPO_CANDIDATE: 40;" in text
    assert "| Rows in the candidate file, not confirmed (C6) | 540 |" in text
    assert (tmp_path / "out/c7_candidates/c7_provider_scores.csv").exists()
    names = {p.name for p in (tmp_path / "out/deliverables").iterdir()}
    assert {"1_provider_rate_file.xlsx", "2_logic_file.md", "3_file_quality_metrics.md", "4_summary.md",
            "7b_cost_effectiveness_metrics.md"} <= names
    logic = (tmp_path / "out/deliverables/2_logic_file.md").read_text()
    assert "Part A: identifying the PPO network" in logic and "Part B: linking TIN -> NPI -> address" in logic
    assert "capitation" in logic and "## 1. Where each link comes from" in logic       # C5's logic, one level down
    quality = (tmp_path / "out/deliverables/3_file_quality_metrics.md").read_text()
    assert all(f"`{m}`" in quality for m in ("schema_validity", "code_coverage", "zero_rates", "staleness"))
    assert "## Scores in this run" in quality
    assert "% TINs linked to a verified address (C5)" in text
    by_net = list(csv.DictReader(open(tmp_path / "out/c3_by_network/c3_file_scores.csv")))
    assert {"ALL", "CONFIRMED_PPO", "CONFIRMED_NON_PPO", "UNRESOLVED"} <= {r["bucket"] for r in by_net}


# ------------------------------------------------------------------ public-source adapters (offline)
def _serve(routes):
    class H(http.server.BaseHTTPRequestHandler):
        def do_GET(self):
            u = urlparse(self.path)
            body = routes(u.path, parse_qs(u.query), f"http://127.0.0.1:{self.server.server_address[1]}")
            data = json.dumps(body).encode()
            self.send_response(200 if body is not None else 404)
            self.send_header("Content-Type", "application/json")
            self.end_headers()
            self.wfile.write(data)

        def log_message(self, *a):
            pass
    srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    return srv, f"http://127.0.0.1:{srv.server_address[1]}"


def test_fhir_plannet_crawler(tmp_path):
    from ptd_c2.config import load_config

    def routes(path, q, base):
        if path.endswith("/Organization"):
            return {"entry": [{"resource": {"resourceType": "Organization", "id": "n1", "name": "PPO - Large Group"}},
                              {"resource": {"resourceType": "Organization", "id": "n2", "name": "HMO - WholeCare"}}]}
        if path.endswith("/PractitionerRole"):
            net = q["network"][0].split("/")[-1]
            if "page" not in q:
                return {"entry": [
                    {"resource": {"resourceType": "PractitionerRole", "practitioner": {"reference": "Practitioner/p1"},
                                  "location": [{"reference": "Location/l1"}], "meta": {"lastUpdated": "2026-09-01T00:00:00Z"}}},
                    {"resource": {"resourceType": "Practitioner", "id": "p1", "identifier": [
                        {"system": "http://hl7.org/fhir/sid/us-npi", "value": "1234567893" if net == "n1" else "1000000004"}]}},
                    {"resource": {"resourceType": "Location", "id": "l1", "address": {
                        "line": ["455 N Brand Blvd"], "city": "Glendale", "state": "CA", "postalCode": "91203"}}}],
                    "link": [{"relation": "next", "url": f"{base}/fhir/PractitionerRole?network={q['network'][0]}&page=2"}]}
            return {"entry": [{"resource": {"resourceType": "PractitionerRole", "practitioner": {"reference": "Practitioner/x"},
                                            "location": [{"reference": "Location/l9"}]}},
                              {"resource": {"resourceType": "Practitioner", "id": "x", "identifier": [
                                  {"system": "http://hl7.org/fhir/sid/us-npi", "value": "1999999999"}]}},
                              {"resource": {"resourceType": "Location", "id": "l9", "address": {
                                  "line": ["1 Elm"], "city": "Reno", "state": "NV", "postalCode": "89501"}}}]}
        if path.endswith("/OrganizationAffiliation"):
            return {"entry": []}
        return None
    srv, base = _serve(routes)
    try:
        n = crawl_plannet(base + "/fhir", tmp_path / "dir.csv", load_config(), log=lambda m: None)
    finally:
        srv.shutdown()
    got = list(csv.DictReader(open(tmp_path / "dir.csv")))
    assert n == 2 and {r["network_label"] for r in got} == {"PPO_FULL", "HMO"}
    assert all(r["state"] == "CA" for r in got)                     # the Nevada location on page 2 is dropped


def test_npi_registry_api_fallback(world, tmp_path):
    import shutil
    n = world["special"]["not_in_nppes"]

    def routes(path, q, base):
        return {"results": [{"enumeration_type": "NPI-1", "basic": {"last_updated": "2026-08-01", "first_name": "A",
                                                                     "last_name": "B"},
                             "taxonomies": [{"code": ca_world.TAX_CARD, "primary": True}],
                             "addresses": [{"address_purpose": "LOCATION", "address_1": "455 N Brand Blvd",
                                            "city": "Glendale", "state": "CA", "postal_code": "91203"}]}]} \
            if q.get("number") == [str(n)] else {"results": []}
    srv, base = _serve(routes)
    db = tmp_path / "r.sqlite"
    shutil.copy(world["db"], db)
    cfg = tmp_path / "c5.yaml"
    text = (Path(__file__).resolve().parent.parent / "config/c5.yaml").read_text()
    cfg.write_text(text.replace("https://npiregistry.cms.hhs.gov/api/", base + "/api/").replace("requests_per_second: 8",
                                                                                               "requests_per_second: 1000")
               .replace("cache_path: data/npi_api_cache.jsonl", f"cache_path: {tmp_path / 'api.jsonl'}"))
    try:
        assert c5(["run", "--db", str(db), "--nppes", str(world["nppes"]), "--npi-api", "10", "--config", str(cfg),
                   "--out", str(tmp_path / "c5"), "--as-of", AS_OF]) == 0
    finally:
        srv.shutdown()
    con = sqlite3.connect(db)
    assert con.execute("SELECT primary_taxonomy, source FROM c5_providers WHERE npi = ?", (n,)).fetchone() == \
        (ca_world.TAX_CARD, "npi_api")


def test_npi_api_replaces_a_missing_nppes_file(world, tmp_path, monkeypatch):
    """No NPPES file: auto mode asks the API for every NPI, gets taxonomies and addresses, and a rerun
    is served from the cache."""
    import shutil
    calls = []

    def routes(path, q, base):
        n = int(q["number"][0])
        calls.append(n)
        if len(calls) == 1:
            return {"results": []}  # an unknown NPI is cached as not found too
        return {"results": [{"enumeration_type": "NPI-1", "basic": {"last_updated": "2026-08-01", "first_name": "A",
                                                                     "last_name": str(n)},
                             "taxonomies": [{"code": ca_world.TAX_CARD, "primary": True}],
                             "addresses": [{"address_purpose": "LOCATION", "address_1": "455 N Brand Blvd",
                                            "city": "Glendale", "state": "CA", "postal_code": "91203"}]}]}
    srv, base = _serve(routes)
    db = tmp_path / "r.sqlite"
    shutil.copy(world["db"], db)
    cfg = tmp_path / "c5.yaml"
    text = (Path(__file__).resolve().parent.parent / "config/c5.yaml").read_text()
    cfg.write_text(text.replace("https://npiregistry.cms.hhs.gov/api/", base + "/api/")
                   .replace("requests_per_second: 8", "requests_per_second: 1000")
                   .replace("cache_path: data/npi_api_cache.jsonl", f"cache_path: {tmp_path / 'api.jsonl'}"))
    monkeypatch.delenv("PTD_NPI_API", raising=False)
    args = ["run", "--db", str(db), "--config", str(cfg), "--out", str(tmp_path / "c5"), "--as-of", AS_OF]
    try:
        assert c5(args) == 0
        con = sqlite3.connect(db)
        npis = con.execute("SELECT COUNT(DISTINCT npi) FROM c4_pairs").fetchone()[0]
        assert sorted(calls) == sorted(set(calls)) and len(calls) >= npis - 5     # each NPI asked once
        found = con.execute("SELECT COUNT(*) FROM c5_providers WHERE source = 'npi_api' AND primary_taxonomy = ?",
                            (ca_world.TAX_CARD,)).fetchone()[0]
        assert found == len(calls) - 1
        assert con.execute("SELECT COUNT(*) FROM c5_addresses WHERE address != ''").fetchone()[0] > 0
        con.close()
        before = len(calls)
        assert c5(args) == 0
        assert len(calls) == before                                              # all from the cache
        summary = json.loads((tmp_path / "c5" / "c5_run_summary.json").read_text())
        assert "npi_api" in summary["sources_used"]
    finally:
        srv.shutdown()


def test_extract_nppes_state_from_monthly_zip(world, tmp_path):
    z = tmp_path / "NPPES_Data_Dissemination_October_2026.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.write(world["nppes"], "npidata_pfile_20050523-20261012.csv")
        zf.write(world["nppes"], "npidata_pfile_20050523-20261012_fileheader.csv")
        zf.write(world["pl"], "pl_pfile_20050523-20261012.csv")
    main, pl = extract_nppes(z, "CA", tmp_path / "ext", log=lambda m: None)
    n_src = sum(1 for _ in open(world["nppes"])) - 1
    assert sum(1 for _ in open(main)) - 1 == n_src and pl.exists()
    assert extract_nppes(z, "NV", tmp_path / "nv", log=lambda m: None)[0].read_text().count("\n") == 1


def test_network_check_sample_and_precision(world, tmp_path):
    sample = world["out"] / "c6/c6_network_check_sample.csv"
    rows_ = list(csv.DictReader(open(sample)))
    assert {r["network_status"] for r in rows_} == {"CONFIRMED_PPO", "PPO_CANDIDATE", "UNRESOLVED"}
    assert all(r["name"] or r["address"] for r in rows_ if r["network_status"] == "CONFIRMED_PPO")
    truth = world["truth"]["ref_network"]
    for r in rows_:  # play the hand-checker with the known answers
        r["directory_result"] = "PPO" if truth.get(int(r["tin"]) - 950000000) == "PPO" else "NOT_PPO"
    filled = tmp_path / "filled.csv"
    with open(filled, "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows_[0]))
        w.writeheader()
        w.writerows(rows_)
    assert c6(["check-network", "--sample", str(filled), "--db", world["db"], "--out", str(tmp_path)]) == 0
    rep = json.loads((tmp_path / "c6_network_accuracy.json").read_text())
    conf = rep["checked"]["by_status"]["CONFIRMED_PPO"]
    assert conf["share_ppo"] == 1.0 and conf["checked"] == min(30, conf["checked"]) and conf["ci95"][0] > 0.85
    text = (tmp_path / "c6_network_accuracy.md").read_text()
    assert "| CONFIRMED_PPO |" in text and "Provider references:" in text


# ------------------------------------------------------------------ PPO directory: capture -> replay -> build -> C4
def _browser():
    import glob as _glob
    import importlib.util
    import os
    if importlib.util.find_spec("playwright") is None:
        return None
    for p in [os.environ.get("PTD_TEST_BROWSER", "")] + _glob.glob("/opt/pw-browsers/chromium-*/chrome-linux/chrome"):
        if p and Path(p).exists():
            return p
    return None


def _directory_site(listings):
    """A tiny directory app: the page POSTs a search; the API answers per ZIP, two providers a page."""
    html = b"""<html><body>search<script>
      fetch('/api/search', {method: 'POST', headers: {'Content-Type': 'application/json'},
        body: JSON.stringify({criteria: {zip: '%s', network: 'PPO'}, paging: {page: 1, size: 2}})})
        .then(r => r.json()).then(d => { document.body.innerText = 'results ' + d.data.results.length; });
    </script></body></html>"""

    class H(http.server.BaseHTTPRequestHandler):
        def do_GET(self):
            self.send_response(200)
            self.send_header("Content-Type", "text/html")
            self.end_headers()
            self.wfile.write(html % listings[0]["zip"].encode())

        def do_POST(self):
            q = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
            z, page = q["criteria"]["zip"], int(q["paging"]["page"])
            hits = [l for l in listings if l["zip"] == z][(page - 1) * 2: page * 2]
            body = {"data": {"results": [
                {**({"npiNumber": str(l["npi"])} if l.get("show_npi") else {}),
                 "provider": {"name": {"first": l["first"], "last": l["last"] + ", MD"}},
                 "locations": [{"name": "Clinic", "address": {"street1": l["street"], "city": l["city"],
                                                              "state": "CA", "zip": l["zip"] + "-0001"}}],
                 "networks": [{"name": "Health Net PPO"}]} for l in hits]}}
            data = json.dumps(body).encode()
            self.send_response(200)
            self.send_header("Content-Type", "application/json")
            self.end_headers()
            self.wfile.write(data)

        def log_message(self, *a):
            pass
    srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    return srv, f"http://127.0.0.1:{srv.server_address[1]}"


def test_directory_capture_replay_build_confirms_ppo(world, tmp_path):
    from ptd_dir.cli import main as ptd_dir
    browser = _browser()
    if browser is None:
        pytest.skip("playwright / a Chromium browser is not available")
    con, truth = world["con"], world["truth"]
    names = dict(rows(con, "SELECT npi, name FROM c5_providers"))
    listings = []
    for g, mode in ((37, "npi"), (38, "name+address"), (39, "name+zip")):
        tin = str(950000000 + g)
        site = truth["tin_site"][tin]
        n = sorted(m for m in truth["tin_npis"][tin] if names.get(m))[0]
        first, last = names[n].split()[0], names[n].split()[-1]
        street = site[0] if mode != "name+zip" else "99 Somewhere Else"
        listings.append({"npi": n, "show_npi": mode == "npi", "first": first, "last": last, "street": street,
                         "city": site[2], "zip": site[3], "mode": mode, "g": g})
    srv, base = _directory_site(listings)
    try:
        cap = tmp_path / "capture.jsonl"
        assert ptd_dir(["capture", "--url", base + "/", "--network", "Health Net PPO", "--out", str(cap),
                        "--browser", browser, "--headless", "--stop-after", "3"]) == 0
        assert cap.exists() and 'criteria' in cap.read_text()
        zips = tmp_path / "zips.txt"
        zips.write_text("\n".join(sorted({l["zip"] for l in listings})) + "\n")
        rep = tmp_path / "replay.jsonl"
        assert ptd_dir(["replay", "--capture", str(cap), "--vary", f"zip=@{zips}", "--vary", "page=1..5",
                        "--delay", "0", "--out", str(rep)]) == 0
        out = tmp_path / "directory.csv"
        assert ptd_dir(["build", "--capture", str(cap), str(rep), "--db", world["db"], "--out", str(out)]) == 0
    finally:
        srv.shutdown()
    built = {int(r["npi"]): r for r in csv.DictReader(open(out))}
    for l in listings:
        assert built[l["npi"]]["npi_source"] == ("listed" if l["mode"] == "npi" else l["mode"])
        assert built[l["npi"]]["network_name"] == "Health Net PPO"
    # C4 with only this directory: listed and name+address confirm, name+zip stays a candidate
    db = str(tmp_path / "r.sqlite")
    assert c4(["run", "--files", str(world["rate_file"]), "--workbook", str(world["workbook"]), "--directory", str(out),
               "--db", db, "--out", str(tmp_path / "c4"), "--workers", "1", "--as-of", AS_OF]) == 0
    st = dict(sqlite3.connect(db).execute("SELECT ref_id, status FROM c4_refs WHERE ref_id IN ('37', '38', '39')"))
    assert st == {"37": "CONFIRMED_PPO", "38": "CONFIRMED_PPO", "39": "PPO_CANDIDATE"}


def test_directory_parser_finds_providers_in_any_shape():
    from ptd_dir.match import name_key
    from ptd_dir.parse import providers_in
    n1, n2 = ca_world.npi(100000077), ca_world.npi(100000078)
    body = {"data": {"results": [
        {"npiNumber": str(n1), "provider": {"name": {"first": "Pat", "last": "Doe7"},
                                            "specialties": [{"description": "Family Medicine"}]},
         "locations": [{"name": "Main Clinic", "address": {"street1": "1200 West Third Street", "city": "Los Angeles",
                                                           "zip": "90017-1234"}}],
         "networks": [{"name": "Health Net PPO"}]},
        {"provider": {"fullName": "Group X"}, "practitioners": [{"npi": str(n2), "firstName": "Bo", "lastName": "Z"}],
         "address": {"line1": "1 Medical Plaza Dr", "city": "Irvine", "zip": "92618"}}]}}
    got = {(r["npi"], r["name"]): r for r in providers_in(body)}
    a = got[(str(n1), "Pat Doe7")]
    assert (a["zip"], a["specialty"], a["listing_networks"]) == ("90017", "Family Medicine", "Health Net PPO")
    assert got[(str(n2), "Bo Z")]["address_line1"] == "1 Medical Plaza Dr"   # practitioner takes the group's site
    assert ("", "Main Clinic") not in got                                     # a location is not a provider
    assert name_key("", "", "LEE, ANN M MD") == name_key("Ann", "Lee, MD") == ("ANN", "LEE")


def test_allowed_amount_file_of_ppo_plans_is_non_ppo_evidence(world, tmp_path):
    n46 = sorted(world["truth"]["tin_npis"]["950000046"])[0]     # HMO ref 46: only inferred non-PPO so far
    oon = tmp_path / "oon.json"
    oon.write_text(json.dumps({"reporting_entity_name": "x", "out_of_network": [
        {"name": "visit", "billing_code_type": "CPT", "billing_code": "99213", "allowed_amounts": [
            {"tin": {"type": "ein", "value": "950000046"}, "billing_class": "professional", "payments": [
                {"allowed_amount": 70.0, "providers": [{"billed_charge": 150.0, "npi": [n46]}]}]}]}]}))
    db = str(tmp_path / "r.sqlite")
    args = ["run", "--files", str(world["rate_file"]), "--workbook", str(world["workbook"]), "--db", db,
            "--out", str(tmp_path / "c4"), "--workers", "1", "--as-of", AS_OF, "--oon", str(oon)]
    assert c4(args + ["--oon-ppo-only"]) == 0
    st, basis = sqlite3.connect(db).execute("SELECT status, status_basis FROM c4_refs WHERE ref_id = '46'").fetchone()
    assert st == "CONFIRMED_NON_PPO" and "oon_ppo_plan" in basis
    assert c4(args) == 0                                           # a shared file is only a note
    st, ev = sqlite3.connect(db).execute("SELECT status, evidence FROM c4_refs WHERE ref_id = '46'").fetchone()
    assert st != "CONFIRMED_NON_PPO" and "shared allowed-amount file" in ev


def _pces_entity(npi, tin, name, street, city, zip5, cycles):
    """A search result shaped like Health Net Find a Provider's (external-api.../pces/query)."""
    return {"carrierEntity": {
        "providerName": name, "providerType": "HOSPITAL", "taxIdNumber": f"{tin[:2]}-{tin[2:]}",
        "nationalProviderIdentifiers": [{"number": str(npi), "type": "Type 2"}],
        "locations": [{"address": {"addressLine1": street, "city": city, "stateCode": "CA", "zipCode": zip5},
                       "specialties": [{"descriptionDisplay": "General Acute Care Hospital"}]}],
        "network": {"product": {"code": "PPO", "name": "HealthNet Commercial PPO", "networkId": "8201"},
                    "carrier": {"code": "PP", "name": "CA PPO Large and Small", "networkId": "8242"},
                    "networkCycles": [{"status": s, "startDate": a + "T00:00:00Z", "endDate": b + "T00:00:00Z"}
                                      for s, a, b in cycles]}},
        "matchedLists": []}


def test_healthnet_find_a_provider_parser_keeps_par_listings_only():
    from ptd_dir.healthnet import is_pces, parse_pces
    body = {"totalHits": 3, "results": [
        _pces_entity(1417554528, "851644866", "USC Norris", "1441 Eastlake Ave", "Los Angeles", "90089",
                     [("Par", "2021-01-01", "4000-01-01")]),
        _pces_entity(1093892275, "941623075", "Siskiyou", "444 Bruce St", "Yreka", "96097",
                     [("Nonpar", "2018-01-01", "2022-12-31"), ("Par", "2023-01-01", "4000-01-01")]),
        _pces_entity(1861409823, "951915820", "Old", "1 Main St", "El Centro", "92243",
                     [("Par", "2015-01-01", "2024-12-31"), ("Nonpar", "2025-01-01", "4000-01-01")])]}
    assert is_pces(body) and not is_pces({"results": [{"type": "name", "results": []}]})
    rows = {r["npi"]: r for r in parse_pces(body, 2026)}
    assert set(rows) == {"1417554528", "1093892275"}                  # Nonpar in 2026 is not in network
    r = rows["1417554528"]
    assert (r["tin"], r["zip"], r["specialty"]) == ("851644866", "90089", "General Acute Care Hospital")
    assert r["listing_networks"] == "HealthNet Commercial PPO | CA PPO Large and Small"
    assert "1861409823" in {r["npi"] for r in parse_pces(body, 2024)}


def test_healthnet_check_captures_build_directory_that_confirms_ppo(world, tmp_path):
    """Captures in `ptd_dir check` format -> build -> C4: an NPI listed Par under the Group PPO network
    confirms its reference; a Nonpar listing does not; a facility listed under another NPI with the same
    TIN is only a candidate."""
    from ptd_dir.cli import main as ptd_dir
    truth = world["truth"]
    net = "PPO - Large Group/Small Group (Platinum, Gold, Silver, Bronze, GF Plans, HSA Compatible, HDHP)"
    ents = []
    for g, cycles, other_npi in ((37, [("Par", "2021-01-01", "4000-01-01")], False),
                                 (38, [("Nonpar", "2021-01-01", "4000-01-01")], False),
                                 (39, [("Par", "2021-01-01", "4000-01-01")], True)):
        tin = str(950000000 + g)
        site = truth["tin_site"][tin]
        npi = ca_world.npi(100000990) if other_npi else sorted(truth["tin_npis"][tin])[0]
        ents.append(_pces_entity(npi, tin, f"Group {g}", site[0], site[2], site[3], cycles))
    cap = tmp_path / "hn_check.jsonl"
    cap.write_text("\n".join(json.dumps({"url": "https://external-api.search.my.centene.com/pces/query", "status": 200,
                                         "network_name": net, "network_year": "2026", "body": {"results": [e]}})
                             for e in ents) + "\n")
    cap.with_suffix(".coverage.jsonl").write_text(json.dumps({"npi": "1", "status": "found"}) + "\n")
    out = tmp_path / "directory.csv"
    assert ptd_dir(["build", "--capture", str(tmp_path / "hn_check*.jsonl"), "--out", str(out)]) == 0  # glob skips the coverage log
    built = list(csv.DictReader(open(out)))
    assert len(built) == 2 and all(r["network_name"] == net for r in built)
    db = str(tmp_path / "r.sqlite")
    assert c4(["run", "--files", str(world["rate_file"]), "--workbook", str(world["workbook"]), "--directory", str(out),
               "--db", db, "--out", str(tmp_path / "c4"), "--workers", "1", "--as-of", AS_OF]) == 0
    st = dict(sqlite3.connect(db).execute("SELECT ref_id, status FROM c4_refs WHERE ref_id IN ('37', '38', '39')"))
    assert st["37"] == "CONFIRMED_PPO"
    assert (st["38"], st["39"]) == ("UNRESOLVED", "PPO_CANDIDATE")


def test_directory_matches_dashed_rate_file_tins():
    from ptd_c2.config import load_config
    from ptd_c4.directory import Directory
    d = Directory([{"npi": 1417554528, "tin": "851644866", "network_name": "PPO - Large Group/Small Group"}],
                  load_config(), ["PPO_FULL"])
    assert d.match(1417554528, "85-1644866") == ({"PPO"}, "npi")             # same TIN, written with a dash
    assert d.match(1417554528, "95-0000000") == ({"PPO"}, "npi_tin_mismatch")
    assert d.match(1999999999, "85-1644866") == ({"PPO"}, "tin")


def test_c4_npi_limit_writes_ppo_check_table(world, tmp_path):
    """C4 --npis keeps only those NPIs' pairs (network tagging still sees the whole file) and writes
    c4_ppo_check.csv: rate-file side + what the Health Net site said, with the name/TIN fallback."""
    from ptd_c4.cli import read_npi_list
    truth = world["truth"]
    t37, t39 = sorted(truth["tin_npis"]["950000037"])[0], sorted(truth["tin_npis"]["950000039"])[0]
    other = ca_world.npi(100000991)
    lim = tmp_path / "ppo_check_results.csv"
    lim.write_text(f"npi,verdict\n{t37},x\n{t39},x\n")
    assert read_npi_list([str(lim), "1234567893, 1417554528"]) == [str(t37), str(t39), "1234567893", "1417554528"]
    net = "PPO - Large Group/Small Group (Platinum, Gold, Silver, Bronze, GF Plans, HSA Compatible, HDHP)"
    s37, s39 = truth["tin_site"]["950000037"], truth["tin_site"]["950000039"]
    recs = [{"status": 200, "network_name": net, "network_year": "2026", "query_kind": "npi",
             "body": {"results": [_pces_entity(t37, "950000037", "Group 37", s37[0], s37[2], s37[3],
                                               [("Par", "2021-01-01", "4000-01-01")])]}},
            {"status": 200, "network_name": net, "network_year": "2026", "query_kind": "name", "query_tin": "950000039",
             "body": {"results": [_pces_entity(other, "950000039", "Group 39 Main", s39[0], s39[2], s39[3],
                                               [("Par", "2021-01-01", "4000-01-01")]),
                                  _pces_entity(ca_world.npi(100000992), "880000001", "Namesake", "1 Elm St", "Fresno",
                                               "93701", [("Par", "2021-01-01", "4000-01-01")])]}}]
    cap = tmp_path / "hn.jsonl"
    cap.write_text("\n".join(json.dumps(r) for r in recs) + "\n")
    cap.with_suffix(".coverage.jsonl").write_text(
        json.dumps({"npi": str(t37), "status": "found", "location": "Los Angeles, CA", "network": net, "year": "2026"}) + "\n"
        + json.dumps({"npi": str(t39), "status": "not_found", "location": "Irvine, CA", "network": net, "year": "2026",
                      "fallback_term": "Group 39", "fallback_same_tin": 1}) + "\n")
    from ptd_dir.cli import main as ptd_dir
    d = tmp_path / "directory.csv"
    assert ptd_dir(["build", "--capture", str(cap), "--out", str(d)]) == 0
    assert {r["tin"] for r in csv.DictReader(open(d))} == {"950000037", "950000039"}   # the namesake is dropped
    db, out = str(tmp_path / "r.sqlite"), tmp_path / "c4"
    assert c4(["run", "--files", str(world["rate_file"]), "--workbook", str(world["workbook"]), "--directory", str(d),
               "--directory-checks", str(cap.with_suffix(".coverage.jsonl")), "--npis", str(lim), "--db", db,
               "--out", str(out), "--workers", "1", "--as-of", AS_OF]) == 0
    con = sqlite3.connect(db)
    assert {n for (n,) in con.execute("SELECT DISTINCT npi FROM c4_pairs")} == {t37, t39}
    assert con.execute("SELECT COUNT(DISTINCT ref_id) FROM c4_refs").fetchone()[0] > 2      # all references tagged
    rows = {int(r["npi"]): r for r in csv.DictReader(open(out / "c4_ppo_check.csv"))}
    assert set(rows) == {t37, t39}
    a, b = rows[t37], rows[t39]
    assert (a["verdict"], a["pair_status"], a["search_location"]) == ("FOUND_PPO_BY_NPI", "CONFIRMED_PPO", "Los Angeles, CA")
    assert a["directory_listing"] == "HealthNet Commercial PPO | CA PPO Large and Small" and int(a["rates_extracted"]) > 0
    assert (b["verdict"], b["npi_search_result"], b["tin_fallback_term"]) == ("NOT_FOUND_BY_NPI_SAME_TIN", "not_found", "Group 39")
    assert b["tin_fallback_npis"] == str(other)


def test_reference_members_keep_the_tin_business_name():
    from ptd_c4.ingest import _group_rows
    rows = _group_rows(1, "7", [{"npi": [1417554528], "tin": {"type": "ein", "value": "85-1644866",
                                                             "business_name": "Keck Medical Center of USC"}}])
    assert rows == [(1, "7", 1417554528, "85-1644866", "ein", "Keck Medical Center of USC")]


def test_c5_directory_address_of_the_same_tin_is_boosted_whatever_the_dashes():
    from ptd_c5.config import load_config as c5cfg
    from ptd_c5.resolve import Resolver
    cfg = c5cfg(as_of=AS_OF)
    r = Resolver(cfg, cfg["_as_of"], {}, {})
    ev = {"source": "directory", "p": 0.75, "detail": json.dumps({"tin": "941623075"})}
    assert r._adjust(ev, "94-1623075") > 0.75 > r._adjust(ev, "95-0000000")


def test_c5_same_tin_directory_site_only_corroborates():
    """St Jude bills under 1194926279; the directory lists the hospital under 1891904942 (same TIN).
    That listing supports the NPPES address in the same ZIP, but another campus of the TIN is not added."""
    from ptd_c2.config import load_config
    from ptd_c4.directory import Directory
    from ptd_c5.sources import directory_tin_candidates
    net = "PPO - Large Group/Small Group"
    d = Directory([{"npi": 1891904942, "tin": "951643325", "network_name": net, "address_line1": "101 E Valencia Mesa Dr",
                    "city": "Fullerton", "state": "CA", "zip": "92835"},
                   {"npi": 1891904943, "tin": "951643325", "network_name": net, "address_line1": "1 Other Campus Rd",
                    "city": "Orange", "state": "CA", "zip": "92868"}], load_config(), ["PPO_FULL"])
    nppes = [{"npi": 1194926279, "source": "nppes_practice", "line1": "101 E VALENCIA MESA DR", "zip": "92835"},
             {"npi": 1194926279, "source": "nppes_mailing", "line1": "9 Billing Way", "zip": "92868"}]
    from datetime import date
    got = directory_tin_candidates(d, [(1194926279, "95-1643325")], nppes, date(2026, 10, 17))
    assert [(c["source"], c["zip"]) for c in got] == [("directory_tin", "92835")]


def _fake_site(shard, out, kw):
    """Stands in for the browser: lists every NPI whose last digit is even under the PPO network."""
    import time as _t
    net = "PPO - Large Group/Small Group (Platinum, Gold, Silver, Bronze, GF Plans, HSA Compatible, HDHP)"
    out = Path(out)
    for t in shard:
        found = int(t["npi"]) % 2 == 0
        ent = [_pces_entity(int(t["npi"]), t.get("tin") or "", "Listed", "1 Main St", "Fresno", "93701",
                            [("Par", "2021-01-01", "4000-01-01")])] if found else []
        with open(out, "a") as fh:
            fh.write(json.dumps({"status": 200, "network_name": net, "network_year": "2026", "query_kind": "npi",
                                 "query_npi": t["npi"], "body": {"results": ent}}) + "\n")
        with open(out.with_suffix(".coverage.jsonl"), "a") as fh:
            fh.write(json.dumps({"npi": t["npi"], "tin": t.get("tin"), "status": "found" if found else "not_found",
                                 "location": "Fresno, CA", "network": net, "year": "2026"}) + "\n")
        _t.sleep(0.02)


def test_c4_checks_every_npi_on_the_site_and_saves_the_table_as_it_goes(world, tmp_path, monkeypatch):
    """C4 --check-directory searches every NPI of the file (resumable), rewrites c4_ppo_check.csv every
    --save-every NPIs, and tags the references with what the site lists."""
    from ptd_c4 import dircheck, ppo_check
    monkeypatch.setattr(dircheck, "_worker", _fake_site)
    monkeypatch.setattr(dircheck, "MP_CONTEXT", "fork")
    monkeypatch.setattr(dircheck, "POLL_SECONDS", 0.05)
    saved = []
    real_write = ppo_check.write
    monkeypatch.setattr(ppo_check, "write", lambda path, rows: (saved.append(len(rows)), real_write(path, rows)))
    db, out, base = str(tmp_path / "r.sqlite"), tmp_path / "c4", tmp_path / "dir" / "hn_all"
    args = ["run", "--files", str(world["rate_file"]), "--workbook", str(world["workbook"]), "--db", db,
            "--out", str(out), "--workers", "1", "--as-of", AS_OF, "--check-directory", "--check-out", str(base),
            "--check-workers", "2", "--save-every", "5"]
    assert c4(args) == 0
    con = sqlite3.connect(db)
    all_npis = {str(n) for (n,) in con.execute("SELECT DISTINCT npi FROM c4_ref_members WHERE npi IS NOT NULL")}
    assert dircheck.checked_npis(dircheck.coverage_paths(base)) == all_npis     # every NPI searched once
    assert len(saved) >= 3 and saved == sorted(saved)                            # the table grew as it went
    rows = list(csv.DictReader(open(out / "c4_ppo_check.csv")))
    assert {r["npi"] for r in rows} == all_npis and {r["site_checked"] for r in rows} == {"yes"}
    found = [r for r in rows if int(r["npi"]) % 2 == 0]
    assert found and {r["verdict"] for r in found} <= {"FOUND_PPO_BY_NPI", "FOUND_PPO_NPI_OTHER_TIN"}
    single = [r for r in found if r["ref_count"] == "1" and r["verdict"] == "FOUND_PPO_BY_NPI"]
    assert "CONFIRMED_PPO" in {r["pair_status"] for r in single}                 # tagged with what the site listed
    assert (base.parent / "hn_all_directory.csv").exists()
    lines = sum(len(open(p).read().splitlines()) for p in dircheck.coverage_paths(base))
    saved.clear()
    assert c4(args) == 0                                                        # a rerun searches nothing again
    assert sum(len(open(p).read().splitlines()) for p in dircheck.coverage_paths(base)) == lines


def test_c4_reuses_an_existing_ppo_check_table_instead_of_searching_again(world, tmp_path, monkeypatch):
    """With a c4_ppo_check.csv from an earlier run, C4 takes the site's results from it: the same verdicts
    and network statuses, and no NPI in it is searched again."""
    from ptd_c4 import dircheck
    monkeypatch.setattr(dircheck, "_worker", _fake_site)
    monkeypatch.setattr(dircheck, "MP_CONTEXT", "fork")
    monkeypatch.setattr(dircheck, "POLL_SECONDS", 0.05)
    common = ["run", "--files", str(world["rate_file"]), "--workbook", str(world["workbook"]), "--workers", "1",
              "--as-of", AS_OF]
    first = tmp_path / "first"
    assert c4(common + ["--db", str(tmp_path / "a.sqlite"), "--out", str(first), "--check-directory",
                        "--check-out", str(tmp_path / "d1" / "hn"), "--save-every", "50"]) == 0
    before = {(r["npi"], r["tin"]): r for r in csv.DictReader(open(first / "c4_ppo_check.csv"))}

    def no_search(shard, out, kw):
        raise AssertionError(f"searched {len(shard)} NPIs again")
    monkeypatch.setattr(dircheck, "_worker", no_search)
    second = tmp_path / "second"
    assert c4(common + ["--db", str(tmp_path / "b.sqlite"), "--out", str(second), "--check-directory",
                        "--check-out", str(tmp_path / "d2" / "hn"), "--ppo-check-csv", str(first / "c4_ppo_check.csv")]) == 0
    assert dircheck.coverage_paths(tmp_path / "d2" / "hn") == []                 # nothing searched
    after = {(r["npi"], r["tin"]): r for r in csv.DictReader(open(second / "c4_ppo_check.csv"))}
    assert set(after) == set(before)
    for k in ("verdict", "pair_status", "npi_search_result", "search_location", "directory_listing"):
        assert {p: r[k] for p, r in after.items()} == {p: r[k] for p, r in before.items()}, k
    assert (second / "c4_directory_used.csv").exists()
    # the default: this output folder's own table, even without --check-directory
    assert c4(common + ["--db", str(tmp_path / "b.sqlite"), "--out", str(second)]) == 0
    again = {(r["npi"], r["tin"]): r["pair_status"] for r in csv.DictReader(open(second / "c4_ppo_check.csv"))}
    assert again == {p: r["pair_status"] for p, r in before.items()}
