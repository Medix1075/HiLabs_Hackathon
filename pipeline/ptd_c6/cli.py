"""Command line entry point.

  python -m ptd_c6 run --workbook "Rate Benchmarking Individual Weights_Hackathon.xlsx"   # -> out/c6/
  python -m ptd_c6 run --workbook ... --nucc nucc_taxonomy_251.csv
  python -m ptd_c6 check-network --sample out/c6/c6_network_check_sample.csv   # after filling it in by hand

Reads the rate store (C4 rates, C5 addresses and provider facts) and writes deliverable 1:
one row per NPI x TIN x address x billing code with the specialty and the headline rate.
"""
from __future__ import annotations

import argparse
import csv
import gzip
import json
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

import yaml

from ptd_c2.cli import write_csv
from ptd_c4.db import DEFAULT_DB, connect, table_exists

from . import __version__
from .network_check import coverage, make_sample, score_sample, write_report, write_sample
from .specialty import assign, load_specialties

DEFAULT_CONFIG = Path(__file__).resolve().parent.parent / "config" / "c6.yaml"
COLUMNS = [
    ("NPI", "npi"), ("TIN", "tin"), ("Address", "address"), ("Billing Code", "code"), ("Specialty", "specialty"),
    ("Rate", "rate"), ("Billing Code Type", "code_type"), ("Code Description", "description"),
    ("NPI Type", "npi_type"), ("Provider Name", "name"), ("Negotiated Type", "negotiated_type"),
    ("Billing Class", "billing_class"), ("Setting", "setting"), ("Place of Service", "pos"),
    ("Modifier", "modifier"), ("Expiration Date", "expiration_date"), ("Price Variants for Code", "variants"),
    ("Address Rank", "address_rank"), ("Address Confidence", "address_confidence"),
    ("Address Verified", "address_verified"), ("Address Sources", "address_sources"), ("City", "city"),
    ("ZIP", "zip5"), ("Region", "region"), ("Specialty Code", "spec_cd"), ("Specialty Source", "specialty_source"),
    ("Code Expected for Specialty", "code_expected"), ("Network Status", "network_status"), ("Network Tag", "network_tag"),
    ("Network Confidence", "network_confidence"), ("Rate Pick Rule", "pick_rule"), ("Source File", "source_file"),
]
C6_FIELDS = ["npi", "tin", "npi_type", "spec_cd", "specialty", "specialty_source", "region", "address_verified",
             "code_type", "code", "negotiated_type", "billing_class", "setting", "pos", "modifier", "rate",
             "code_expected", "network_status", "network_tag", "network_confidence", "file_id"]
ABOUT = {
    "NPI / TIN": "From the rate file's provider_groups (C4). NPPES hides TINs, so the pairing comes only from the rate file.",
    "Address": "C5's resolved practice address. A pair practising at several sites has one row per site (Address Rank). "
               "Blank = no source had an address (see Address Sources / c5_addresses.csv).",
    "Billing Code / Rate": "C4's clean rate. When a code has several price variants (billing class, setting, place of "
                           "service, modifier, rate type) the headline one is shown; Price Variants for Code says how "
                           "many exist and the c4_rates view in out/rates.sqlite has them all.",
    "Specialty": "C6: NPPES taxonomy -> HiLabs crosswalk, else Medicare specialty, else inferred from billed codes, "
                 "else NUCC. Specialty Source says which.",
    "Code Expected for Specialty": "TRUE if the code is in the specialty's HiLabs code list, FALSE if not (worth a look), "
                                   "blank if the specialty has no code list.",
    "Network Status": "C4: CONFIRMED_PPO (direct evidence: a PPO directory listing that can only be about this "
                      "contract, the reference's own network name, or a PPO-only file). The candidate file holds "
                      "PPO_CANDIDATE (inferred or weak evidence) and UNRESOLVED rows; see c4_network_tags.csv for the "
                      "evidence behind each provider reference.",
    "Network Tag": "C4's evidence vote (PPO / NON_PPO / UNKNOWN) behind the status, with its confidence.",
}
OUTPUTS = {  # file stem -> which network statuses go in it (config/c6.yaml `network`)
    "c6_provider_rate_file": "rate_file_statuses",
    "c6_provider_rate_file_candidates": "candidate_file_statuses",
}


def load_config(path=None) -> dict:
    with open(path or DEFAULT_CONFIG, encoding="utf-8") as fh:
        cfg = yaml.safe_load(fh)
    cfg["_path"] = str(path or DEFAULT_CONFIG)
    return cfg


def headline_key(v: dict, npi_type: str, hc: dict):
    """Sort key: lower is preferred."""
    def idx(order, val):
        return order.index(val) if val in order else len(order)
    classes = hc["billing_class"].get(str(npi_type), hc["billing_class"]["1"])
    pos = v["pos"].split(",") if v["pos"] else []
    return (idx(hc["negotiated_type"], v["negotiated_type"]), idx(classes, v["billing_class"]),
            0 if any(p in pos for p in hc["preferred_pos"]) else (1 if not pos else 2),
            (1 if v["modifier"] else 0) if hc["prefer_no_modifier"] else 0,
            idx(hc["setting"], v["setting"]), v["rate"])


class ExcelWriter:
    """Streaming xlsx writer (openpyxl write-only): new sheet every rows_per_sheet rows, stops at max_rows."""

    def __init__(self, path: Path, cfg: dict, header: list[str] | None = None):
        import openpyxl

        self.path, self.cfg = path, cfg
        self.header = header or [c for c, _ in COLUMNS]
        self.wb = openpyxl.Workbook(write_only=True)
        self.about = self.wb.create_sheet("About")
        self.ws, self.rows, self.sheet_rows, self.sheets, self.truncated = None, 0, 0, 0, False

    def add(self, row: list) -> None:
        if self.rows >= self.cfg["max_rows"]:
            self.truncated = True
            return
        if self.ws is None or self.sheet_rows >= self.cfg["rows_per_sheet"]:
            self.sheets += 1
            self.ws = self.wb.create_sheet("Rates" if self.sheets == 1 else f"Rates {self.sheets}")
            self.ws.append(self.header)
            self.sheet_rows = 0
        self.ws.append(row)
        self.rows += 1
        self.sheet_rows += 1

    def close(self, about_rows: list[list]) -> None:
        if self.ws is None:  # no rows: still give the columns
            self.ws = self.wb.create_sheet("Rates")
            self.ws.append(self.header)
        for r in about_rows:
            self.about.append(r)
        self.wb.save(self.path)


def run(args) -> int:
    cfg = load_config(args.config)
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    log_lines: list[str] = []

    def log(msg: str) -> None:
        print(msg, flush=True)
        log_lines.append(msg)

    started = datetime.now(timezone.utc)
    log(f"C6 v{__version__} started {started.isoformat()} config={cfg['_path']}")
    con = connect(args.db)
    if not table_exists(con, "c4_sig_rates"):
        print(f"No C4 rate store in {args.db}. Run C4 first.", file=sys.stderr)
        return 1
    has_c5 = table_exists(con, "c5_addresses")
    if not has_c5:
        log("[input] no C5 addresses in the rate store: Address columns stay blank (run python -m ptd_c5 run)")
    sp = load_specialties(args.workbook, args.nucc)
    log(f"[input] specialties: {len(sp.by_taxonomy)} taxonomies in the crosswalk, {len(sp.codes)} code profiles, "
        f"{len(sp.nucc)} NUCC codes")

    # ------------------------------------------------------------ providers and specialties
    providers = {}
    if has_c5:
        cols = [d[1] for d in con.execute("PRAGMA table_info(c5_providers)")]
        providers = {r[0]: dict(zip(cols, r)) for r in con.execute("SELECT * FROM c5_providers")}
    # only NPIs that have at least one extracted rate (pairs whose codes are all out of scope have no rows)
    npis = [r[0] for r in con.execute("SELECT DISTINCT npi FROM c4_pairs WHERE sig_id IN "
                                      "(SELECT sig_id FROM c4_sig_rates)")]
    need_codes = []
    for n in npis:
        p = providers.get(n, {})
        taxes = [t for t in str(p.get("taxonomies") or "").split("|") if t]
        if not any(t.upper() in sp.by_taxonomy for t in taxes) and str(p.get("entity_type")) != "2":
            need_codes.append(n)
    codes_by_npi: dict[int, set] = defaultdict(set)
    if need_codes and sp.codes:
        con.execute("DROP TABLE IF EXISTS temp.need")
        con.execute("CREATE TEMP TABLE need (npi INTEGER PRIMARY KEY)")
        con.executemany("INSERT INTO temp.need VALUES (?)", [(n,) for n in need_codes])
        for n, code in con.execute("SELECT DISTINCT p.npi, s.code FROM temp.need t CROSS JOIN c4_pairs p ON p.npi = t.npi "
                                   "CROSS JOIN c4_sig_rates s ON s.file_id = p.file_id AND s.sig_id = p.sig_id"):
            codes_by_npi[n].add(code)
    sp_cfg = cfg["specialty"]
    spec: dict[int, dict] = {}
    for n in npis:
        p = providers.get(n, {})
        taxes = [t for t in str(p.get("taxonomies") or "").split("|") if t]
        codes = codes_by_npi.get(n, set())
        res = assign(p.get("entity_type") or "", p.get("primary_taxonomy") or "", taxes, p.get("dac_specialty") or "",
                     codes if len(codes) >= sp_cfg["infer_min_codes"] else set(), sp)
        spec[n] = {**res, "npi_type": str(p.get("entity_type") or ""), "name": p.get("name") or ""}
    log("[c6] specialty source: " + ", ".join(f"{k} {v:,}" for k, v in
                                              Counter(s["specialty_source"] for s in spec.values()).most_common()))

    # ------------------------------------------------------------ addresses, codes, files
    addresses: dict[tuple, list] = defaultdict(list)
    if has_c5:
        for r in con.execute("SELECT npi, tin, rank, address, city, zip5, region, confidence, verified, sources "
                             "FROM c5_addresses ORDER BY npi, tin, rank"):
            addresses[(r[0], r[1])].append(r[2:])
    descriptions = {(t, c): d for t, c, d in con.execute("SELECT code_type, code, description FROM c4_codes")}
    files = {fid: Path(str(url)).name for fid, url in con.execute("SELECT file_id, url FROM c4_files")}
    pairs_by_sig: dict[tuple, list] = defaultdict(list)
    for fid, sig, npi, tin in con.execute("SELECT file_id, sig_id, npi, tin FROM c4_pairs"):
        pairs_by_sig[(fid, sig)].append((npi, tin))

    con.execute("DROP TABLE IF EXISTS c6_rates")
    con.execute(f"CREATE TABLE c6_rates ({', '.join(C6_FIELDS)})")
    hc = cfg["headline"]
    route: dict[str, str] = {}
    for stem, key in OUTPUTS.items():
        for st in cfg["network"][key]:
            route[st] = stem
    sinks = {stem: {"xl": ExcelWriter(out / f"{stem}.xlsx", cfg["excel"]),
                    "fh": gzip.open(out / f"{stem}.csv.gz", "wt", newline="", encoding="utf-8")} for stem in OUTPUTS}
    for sk in sinks.values():
        sk["csv"] = csv.writer(sk["fh"])
        sk["csv"].writerow([c for c, _ in COLUMNS])
    stats = Counter()
    by_file = {stem: Counter() for stem in OUTPUTS}
    c6_buf: list[tuple] = []
    var_cols = ["code_type", "code", "negotiated_type", "billing_class", "setting", "pos", "modifier", "rate",
                "expiration_date", "network_status", "network_tag", "network_confidence", "pick_rule"]

    def emit_sig(key, variants_by_code: dict) -> None:
        fid, _ = key
        for npi, tin in pairs_by_sig.get(key, ()):
            s = spec[npi]
            ntype = s["npi_type"] if s["npi_type"] in ("1", "2") else "1"
            addrs = addresses.get((npi, tin)) or [(1, "", "", "", "", 0.0, False, "")]
            region = addrs[0][4]
            for (ctype, code), vs in variants_by_code.items():
                v = min(vs, key=lambda x: headline_key(x, ntype, hc))
                expected = sp.code_expected(s["spec_cd"], code) if s["spec_cd"] else None
                c6_buf.append((npi, tin, s["npi_type"], s["spec_cd"], s["specialty"], s["specialty_source"], region,
                               bool(addrs[0][6]), ctype, code, v["negotiated_type"], v["billing_class"], v["setting"],
                               v["pos"], v["modifier"], v["rate"], expected, v["network_status"], v["network_tag"],
                               v["network_confidence"], fid))
                stats["pair_codes"] += 1
                for rank, address, city, zip5, reg, conf, verified, srcs in addrs:
                    row = {**v, "npi": npi, "tin": tin, "address": address, "specialty": s["specialty"],
                           "description": descriptions.get((ctype, code), ""), "npi_type": s["npi_type"],
                           "name": s["name"], "variants": len(vs), "address_rank": rank,
                           "address_confidence": conf, "address_verified": bool(verified), "address_sources": srcs,
                           "city": city, "zip5": zip5, "region": reg, "spec_cd": s["spec_cd"],
                           "specialty_source": s["specialty_source"], "code_expected": expected,
                           "source_file": files.get(fid, "")}
                    cells = [row.get(k, "") for _, k in COLUMNS]
                    stats["rows"] += 1
                    stem = route.get(v["network_status"])
                    if stem is None:
                        stats["rows_not_written"] += 1
                        continue
                    sinks[stem]["csv"].writerow(cells)
                    sinks[stem]["xl"].add(cells)
                    by_file[stem]["rows"] += 1
                    by_file[stem]["rows_with_address"] += bool(address)
                    stats["rows_with_address"] += bool(address)
                    stats["rows_code_unexpected"] += expected is False
        if len(c6_buf) >= 50000:
            con.executemany(f"INSERT INTO c6_rates VALUES ({','.join('?' * len(C6_FIELDS))})", c6_buf)
            c6_buf.clear()

    cur, by_code = None, defaultdict(list)
    for r in con.execute(f"SELECT file_id, sig_id, {', '.join(var_cols)} FROM c4_sig_rates "
                         "ORDER BY file_id, sig_id"):
        key = (r[0], r[1])
        if key != cur:
            if cur is not None:
                emit_sig(cur, by_code)
            cur, by_code = key, defaultdict(list)
        v = dict(zip(var_cols, r[2:]))
        by_code[(v["code_type"], v["code"])].append(v)
    if cur is not None:
        emit_sig(cur, by_code)
    for sk in sinks.values():
        sk["fh"].close()
    if c6_buf:
        con.executemany(f"INSERT INTO c6_rates VALUES ({','.join('?' * len(C6_FIELDS))})", c6_buf)
    con.execute("CREATE INDEX ix_c6_rates ON c6_rates (code_type, code)")
    con.commit()

    summary = {
        "version": __version__, "started": started.isoformat(), "finished": datetime.now(timezone.utc).isoformat(),
        "rows": stats["rows"], "rows_with_address": stats["rows_with_address"], "pair_codes": stats["pair_codes"],
        "rows_code_not_expected_for_specialty": stats["rows_code_unexpected"],
        "npis": len(npis), "specialty_sources": dict(Counter(s["specialty_source"] for s in spec.values())),
        "files": {stem: {"statuses": cfg["network"][OUTPUTS[stem]], "rows": by_file[stem]["rows"],
                         "rows_with_address": by_file[stem]["rows_with_address"], "excel_rows": sinks[stem]["xl"].rows,
                         "excel_sheets": sinks[stem]["xl"].sheets, "excel_truncated": sinks[stem]["xl"].truncated}
                  for stem in OUTPUTS},
        "pair_codes_by_network_status": dict(con.execute("SELECT network_status, COUNT(*) FROM c6_rates GROUP BY 1")),
    }
    titles = {"c6_provider_rate_file": "Provider rate file (deliverable 1): confirmed PPO providers only",
              "c6_provider_rate_file_candidates": "Candidate PPO rates: not confirmed, kept for review (NOT deliverable 1)"}
    for stem, sk in sinks.items():
        xl = sk["xl"]
        about = [[titles[stem], ""], ["Generated", summary["finished"]],
                 ["Network statuses in this file", ", ".join(cfg["network"][OUTPUTS[stem]])],
                 ["Rows", by_file[stem]["rows"]], ["Rows with an address", by_file[stem]["rows_with_address"]],
                 ["Excel complete", f"no: see {stem}.csv.gz" if xl.truncated else "yes"], ["", ""],
                 ["Column", "Meaning"]] + [[k, v] for k, v in ABOUT.items()]
        xl.close(about)
    write_csv(out / "c6_providers.csv",
              [{"npi": n, **s} for n, s in sorted(spec.items())],
              ["npi", "npi_type", "name", "spec_cd", "specialty", "specialty_source", "specialty_note"])
    (out / "c6_run_summary.json").write_text(json.dumps(summary, indent=2, default=str))
    for stem in OUTPUTS:
        f = summary["files"][stem]
        log(f"[c6] {stem}: {f['rows']:,} rows ({f['rows_with_address']:,} with an address), network status "
            f"{', '.join(f['statuses'])}; Excel {f['excel_rows']:,} rows in {f['excel_sheets']} sheet(s)"
            + (" (truncated: the CSV has every row)" if f["excel_truncated"] else ""))
    if not summary["files"]["c6_provider_rate_file"]["rows"]:
        log("[c6] NOTE: no rate has a confirmed PPO status, so the provider rate file is empty. Confirmation needs "
            "target-specific evidence, usually a PPO directory table (python -m ptd_c4 run --directory ...). "
            "The candidate file holds the unconfirmed rates.")
    sample = make_sample(con, cfg["network"]["check_sample_per_status"])
    write_sample(out / "c6_network_check_sample.csv", sample)
    write_report(out / "c6_network_accuracy.md", coverage(con), None, str(out / "c6_network_check_sample.csv"))
    log(f"[c6] network check sample: {len(sample)} NPI-TIN pairs to look up in the PPO directory "
        f"({out / 'c6_network_check_sample.csv'}); coverage in {out / 'c6_network_accuracy.md'}")
    (out / "c6_run_log.txt").write_text("\n".join(log_lines) + "\n")
    con.close()
    log(f"[done] outputs in {out}")
    return 0


def check_network(args) -> int:
    """Score a hand-checked network sample and rewrite the accuracy report."""
    con = connect(args.db)
    checked = score_sample(Path(args.sample))
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    write_report(out / "c6_network_accuracy.md", coverage(con), checked, args.sample)
    con.close()
    print(json.dumps(checked, indent=2))
    return 0


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(prog="ptd_c6", description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    r = sub.add_parser("run", help="write the provider rate file")
    r.add_argument("--db", default=str(DEFAULT_DB))
    r.add_argument("--workbook", help="HiLabs Rate Benchmarking Individual Weights workbook (taxonomy crosswalk, code lists)")
    r.add_argument("--nucc", help="NUCC taxonomy CSV (names for taxonomies the crosswalk does not cover)")
    r.add_argument("--config", help="config YAML (default config/c6.yaml)")
    r.add_argument("--out", default="out/c6")
    r.set_defaults(func=run)
    c = sub.add_parser("check-network", help="score a hand-checked c6_network_check_sample.csv (PPO precision)")
    c.add_argument("--sample", required=True, help="the sample with directory_result filled in")
    c.add_argument("--db", default=str(DEFAULT_DB))
    c.add_argument("--out", default="out/c6")
    c.set_defaults(func=check_network)
    args = ap.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
