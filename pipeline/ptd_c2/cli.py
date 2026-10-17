"""Command line entry point.

  python -m ptd_c2 run --run-c1                     # C1 (scrape + parse) then C2
  python -m ptd_c2 run                              # C2 on C1's outputs in data/
  python -m ptd_c2 run --c1 c1_output.csv --index raw_indexes/ --out out/c2
  python -m ptd_c2 validate --sample out/c2/c2_validation_sample.csv --out out/c2
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from ptd_pipeline.__main__ import main as run_c1
from ptd_pipeline.index_parser import DEFAULT_RAW_DIR as C1_RAW_DIR
from ptd_pipeline.index_parser import DEFAULT_TABLE_CSV as C1_TABLE_CSV

from . import __version__
from .combine import combine, ppo_file_list
from .config import load_config
from .file_vote import vote_files
from .loaders import discover_index_sources, load_c1, load_oon_map
from .name_signals import assign_name_signals
from .network_probe import probe_files
from .scope import apply_scope
from .supporting import oon_signal, payment_signal, probe_oon_files
from .validation import SAMPLE_FIELDS, make_sample, score_sample

CLASSIFICATION_FIELDS = [
    "row_id", "brand", "plan_name", "plan_id", "market_type", "rate_file_url", "state", "state_source",
    "network_names", "name_signal", "name_tokens", "file_vote", "oon_signal", "oon_measured",
    "payment_signal", "payment_measured", "final_network_type", "ppo_variant", "in_scope_ppo",
    "needs_network_filter", "confidence", "decided_by", "reason",
]


def _cell(v):
    if isinstance(v, (dict, list, set, tuple)):
        return json.dumps(sorted(v) if isinstance(v, set) else v, default=str)
    return v


def write_csv(path: Path, rows: list[dict], fields: list[str] | None = None) -> None:
    fields = fields or (list(rows[0].keys()) if rows else ["(empty)"])
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({k: _cell(r.get(k, "")) for k in fields})


def write_xlsx(path: Path, sheets: dict[str, tuple[list[dict], list[str] | None]]) -> None:
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, (rows, fields) in sheets.items():
        ws = wb.create_sheet(name[:31])
        fields = fields or (list(rows[0].keys()) if rows else ["(empty)"])
        ws.append(fields)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([str(_cell(r.get(k, ""))) if isinstance(r.get(k), (dict, list, set)) else r.get(k, "")
                       for k in fields])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def _run_c1_first(args) -> int:
    """Run C1 (Step 0 scrape + Step 1 index parse) so C2 reads this month's files."""
    c1_args = ["--html-file", str(args.c1_html_file)] if args.c1_html_file else []
    print("[c1] running C1: scrape the Centene page, then parse every index file", flush=True)
    code = run_c1(c1_args)
    if code == 1:
        print("[c1] C1 failed (see log above); C2 not run", flush=True)
    elif code == 2:
        print("[c1] some index files failed to parse; C2 continues on the ones that worked", flush=True)
    return code


def run(args) -> int:
    if args.run_c1 and _run_c1_first(args) == 1:
        return 1
    args.c1 = args.c1 or str(C1_TABLE_CSV)
    if not Path(args.c1).exists():
        print(f"C1 table {args.c1} not found. Run C1 first (python -m ptd_pipeline), "
              "add --run-c1, or pass --c1 <table>.", file=sys.stderr)
        return 1
    if args.index is None and C1_RAW_DIR.is_dir():
        args.index = [str(C1_RAW_DIR)]  # C1 caches the raw index files here
    cfg = load_config(args.config)
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    log_lines: list[str] = []

    def log(msg: str) -> None:
        print(msg, flush=True)
        log_lines.append(msg)

    started = datetime.now(timezone.utc)
    log(f"C2 v{__version__} started {started.isoformat()} config={cfg['_path']}")
    for note in cfg.get("context_notes", []):
        log(f"[context] {note}")

    # Input
    rows, c1_stats = load_c1(args.c1, args.sheet)
    log(f"[input] C1 rows {c1_stats['input_rows']}, duplicates dropped {c1_stats['duplicates_dropped']}, "
        f"columns {c1_stats['column_mapping']}")

    # Step 1
    in_scope, exclusions = apply_scope(rows, cfg)
    by_rule = Counter(e["primary_rule"] for e in exclusions)
    trap = sum(1 for e in exclusions if e["name_mentions_ppo"])
    log(f"[step1] kept {len(in_scope)} of {len(rows)} rows; excluded by rule {dict(by_rule)}; "
        f"{trap} excluded rows had 'PPO' in the name (removed by market/brand/state, not keyword)")
    inferred = sum(1 for r in in_scope if r["state_source"].startswith("brand_default"))
    if inferred:
        log(f"[step1] {inferred} kept rows had no explicit state; CA inferred from the Health Net brand")

    # Step 2
    hyp_report = assign_name_signals(in_scope, cfg)
    log(f"[step2] name signals {dict(Counter(r['name_signal'] for r in in_scope))}")
    for h in hyp_report:
        log(f"[step2] hypothesis {h['hypothesis']}: support {h['support_on_labelled_plans']}, "
            f"precision {h['precision']}, {h['verdict']}")

    # Step 3
    votes = vote_files(in_scope, cfg)
    log(f"[step3] {len(votes)} distinct rate files; labels {dict(Counter(v['file_vote'] for v in votes.values()))}")

    # Step 4
    probes: dict = {}
    if cfg["network_probe"]["enabled"] and not args.no_probe:
        if cfg["network_probe"]["probe_scope"] == "vote_candidates":
            urls = [u for u, v in votes.items() if v["file_vote"] != "NON_PPO"]
        else:
            urls = list(votes)
        probes = probe_files(sorted(urls), cfg, out / "network_probe_cache.json", log)
        log(f"[step4] network labels {dict(Counter(p['network_label'] for p in probes.values()))}")
    else:
        log("[step4] SKIPPED (--no-probe or disabled): no network_name confirmation; confidence capped at medium")
    for r in in_scope:
        p = probes.get(r["rate_file_url"])
        r["network_names"] = "|".join(p["network_names"]) if p and p.get("network_names") else ""

    # Step 5
    index_sources = discover_index_sources(args.index or [])
    oon_map, index_log = ({"by_plan_file": {}, "by_plan": {}}, [])
    if index_sources:
        oon_map, index_log = load_oon_map(index_sources, cfg["network_probe"])
        for il in index_log:
            log(f"[step5] index {il['index_source']}: {il['status']}, {il['plans']} plans, "
                f"{il['plans_with_oon_file']} with allowed_amount_file")
    else:
        log("[step5] no raw index files given (--index): oon_signal will be 'unknown'")
    for r in in_scope:
        key3 = (r["plan_name_norm"], r["plan_id"], r["rate_file_url"])
        urls = oon_map["by_plan_file"].get(key3) or oon_map["by_plan"].get(key3[:2]) or set()
        if r.get("oon_file_url"):
            urls = set(urls) | {r["oon_file_url"]}
        r["_oon_urls"] = set(urls)
    oon_enabled = cfg["supporting"]["oon_probe_enabled"] and not args.no_probe
    oon_results = probe_oon_files([u for r in in_scope for u in r["_oon_urls"]], cfg, log) if oon_enabled else {}
    for r in in_scope:
        r["oon_signal"], r["oon_measured"] = oon_signal(r["_oon_urls"], oon_results, oon_enabled)
        r["payment_signal"], r["payment_measured"] = payment_signal(probes.get(r["rate_file_url"]), cfg)
    log(f"[step5] oon {dict(Counter(r['oon_signal'] for r in in_scope))}; "
        f"payment {dict(Counter(r['payment_signal'] for r in in_scope))}")

    # Step 6
    conflicts, support_report = combine(rows, votes, probes, cfg)
    ppo_files = ppo_file_list(rows, probes)
    hyp_report += support_report
    for h in support_report:
        log(f"[step5] hypothesis {h['hypothesis']}: support {h['support_on_labelled_plans']}, "
            f"precision {h['precision']}, {h['verdict']}")
    final_counts = Counter(r["final_network_type"] for r in rows)
    log(f"[step6] final types {dict(final_counts)}; in-scope PPO rows {sum(r['in_scope_ppo'] for r in rows)}; "
        f"confidence {dict(Counter(r['confidence'] for r in rows if r['in_scope']))}; "
        f"conflicts logged {len(conflicts)}; PPO rate files {len(ppo_files)}")
    assert all(r.get("reason") for r in rows), "every row must carry a reason"

    # Step 7
    sample = make_sample(rows, cfg)

    # Outputs
    probe_rows = []
    for url, p in sorted(probes.items()):
        probe_rows.append({k: p.get(k, "") for k in (
            "rate_file_url", "status", "network_label", "network_names", "network_labels", "ppo_network_names",
            "provider_references_read", "provider_references_complete", "in_network_items_read", "arrangements",
            "distinct_tins", "tins_with_50plus_npis", "blank_network_names", "location_refs", "location_files_fetched", "bytes_read",
            "seconds", "stop_reason", "top_level", "note", "error")})
    name_rows = [{k: r[k] for k in ("row_id", "plan_name", "plan_name_norm", "plan_id", "rate_file_url",
                                     "name_signal", "name_types", "name_strength", "name_tokens")} for r in in_scope]
    support_rows = [{"row_id": r["row_id"], "plan_name": r["plan_name"], "rate_file_url": r["rate_file_url"],
                     "oon_file_urls": "|".join(sorted(u for u in r["_oon_urls"] if u)),
                     "oon_signal": r["oon_signal"], "oon_measured": r["oon_measured"],
                     "payment_signal": r["payment_signal"], "payment_measured": r["payment_measured"]}
                    for r in in_scope]
    vote_rows = sorted(votes.values(), key=lambda v: v["rate_file_url"])
    for r in rows:
        r.setdefault("network_names", "")
        for k in ("name_signal", "name_tokens", "file_vote", "oon_signal", "oon_measured",
                  "payment_signal", "payment_measured"):
            r.setdefault(k, "")

    files = {
        "c2_classification.csv": (rows, CLASSIFICATION_FIELDS),
        "c2_ppo_rate_files.csv": (ppo_files, None),
        "c2_exclusions.csv": (exclusions, None),
        "c2_name_signals.csv": (name_rows, None),
        "c2_file_votes.csv": (vote_rows, None),
        "c2_network_probe.csv": (probe_rows, None),
        "c2_supporting.csv": (support_rows, None),
        "c2_conflicts.csv": (conflicts, ["row_id", "plan_name", "plan_id", "rate_file_url", "conflict_type",
                                         "detail", "resolution"]),
        "c2_hypotheses.csv": (hyp_report, None),
        "c2_index_files.csv": (index_log, None),
        "c2_validation_sample.csv": (sample, SAMPLE_FIELDS),
    }
    for name, (data, fields) in files.items():
        write_csv(out / name, data, fields)
    write_xlsx(out / "c2_results.xlsx", {
        "classification": files["c2_classification.csv"], "ppo_rate_files": files["c2_ppo_rate_files.csv"],
        "exclusions": files["c2_exclusions.csv"], "file_votes": files["c2_file_votes.csv"],
        "network_probe": files["c2_network_probe.csv"], "conflicts": files["c2_conflicts.csv"],
        "hypotheses": files["c2_hypotheses.csv"], "validation_sample": files["c2_validation_sample.csv"],
    })
    summary = {
        "version": __version__, "started": started.isoformat(), "finished": datetime.now(timezone.utc).isoformat(),
        "config": cfg["_path"], "c1": str(args.c1), "c1_stats": c1_stats, "index_sources": index_sources,
        "network_probe_ran": bool(probes), "rows": len(rows), "in_scope_rows": len(in_scope),
        "exclusions_by_rule": dict(by_rule), "final_types": dict(final_counts),
        "in_scope_ppo_rows": sum(r["in_scope_ppo"] for r in rows), "ppo_rate_files": len(ppo_files),
        "conflicts": dict(Counter(c["conflict_type"] for c in conflicts)),
        "include_enhancedcare_ppo": cfg["network_probe"]["include_enhancedcare_ppo"],
        "context_notes": cfg.get("context_notes", []),
    }
    (out / "c2_run_summary.json").write_text(json.dumps(summary, indent=2, default=str))
    (out / "c2_run_log.txt").write_text("\n".join(log_lines) + "\n")
    log(f"[done] outputs in {out}")
    return 0


def validate(args) -> int:
    rows, summary = score_sample(args.sample)
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    write_csv(out / "c2_validation_results.csv", rows, SAMPLE_FIELDS + ["match"])
    (out / "c2_validation_summary.json").write_text(json.dumps(summary, indent=2))
    print(json.dumps(summary, indent=2))
    return 0


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(prog="ptd_c2", description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    r = sub.add_parser("run", help="classify plans and produce the PPO rate-file list")
    r.add_argument("--c1", help=f"C1 output table (.csv, .xlsx, .json, .jsonl); default {C1_TABLE_CSV}")
    r.add_argument("--run-c1", action="store_true", help="run C1 (ptd_pipeline) first, then C2 on its outputs")
    r.add_argument("--c1-html-file", type=Path, help="with --run-c1: parse a saved copy of the Centene page")
    r.add_argument("--sheet", help="sheet name when --c1 is an Excel file")
    r.add_argument("--index", nargs="*", help="raw index JSON files, directories or URLs (for allowed_amount_file); "
                                           f"default {C1_RAW_DIR} when it exists")
    r.add_argument("--config", help="config YAML (default config/c2.yaml)")
    r.add_argument("--out", default="out/c2", help="output directory")
    r.add_argument("--no-probe", action="store_true", help="skip reading rate/OON files (offline run)")
    r.set_defaults(func=run)
    v = sub.add_parser("validate", help="score a filled-in validation sample")
    v.add_argument("--sample", required=True)
    v.add_argument("--out", default="out/c2")
    v.set_defaults(func=validate)
    args = ap.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
